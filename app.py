# ============================================================
# FIX: Python/OpenSSL md5 compatibility
# ============================================================

import hashlib

_original_md5 = hashlib.md5


def compatible_md5(data=b'', *args, **kwargs):
    kwargs.pop("usedforsecurity", None)
    return _original_md5(data, *args, **kwargs)


hashlib.md5 = compatible_md5


# ============================================================
# IMPORTS
# ============================================================

from flask import (
    Flask,
    render_template,
    request,
    redirect,
    url_for,
    send_from_directory,
    jsonify
)

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

import os
import tempfile
import sqlite3
import json
import shutil

from datetime import datetime, date

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


# ============================================================
# FLASK CONFIG
# ============================================================

app = Flask(__name__)


BASE_DIR = os.path.dirname(
    os.path.abspath(__file__)
)


DATA_FOLDER = os.path.join(
    BASE_DIR,
    "data"
)


# ============================================================
# FILE PATHS
# ============================================================

STOCK_FILE = os.path.join(
    DATA_FOLDER,
    "stock.xlsx"
)


INTERACTION_FILE = os.path.join(
    DATA_FOLDER,
    "drug_interactions.xlsx"
)


# ไฟล์ต้นฉบับใบสั่งยาที่ระบบเก็บไว้เพื่อ Sync แบบ Real-time
PRESCRIPTION_SOURCE_FILE = os.path.join(
    DATA_FOLDER,
    "prescription_source.xlsx"
)


CONSULT_FOLDER = os.path.join(
    BASE_DIR,
    "static",
    "consults"
)


DATABASE_FILE = os.path.join(
    DATA_FOLDER,
    "medication_history.db"
)


# ============================================================
# CREATE FOLDERS
# ============================================================

os.makedirs(
    DATA_FOLDER,
    exist_ok=True
)


os.makedirs(
    CONSULT_FOLDER,
    exist_ok=True
)


# ============================================================
# DATABASE
# ============================================================

def get_db():

    conn = sqlite3.connect(
        DATABASE_FILE,
        timeout=10
    )

    conn.row_factory = sqlite3.Row

    return conn


def add_column_if_missing(
    conn,
    table_name,
    column_name,
    column_type
):

    columns = conn.execute(
        "PRAGMA table_info({})".format(
            table_name
        )
    ).fetchall()


    existing_columns = [
        row["name"]
        for row in columns
    ]


    if column_name not in existing_columns:

        conn.execute(
            "ALTER TABLE {} ADD COLUMN {} {}".format(
                table_name,
                column_name,
                column_type
            )
        )


# ============================================================
# INIT DATABASE
# ============================================================

def init_database():

    conn = get_db()


    try:

        # ====================================================
        # PRESCRIPTION QUEUE
        # ====================================================

        conn.execute("""
            CREATE TABLE IF NOT EXISTS prescription_queue (

                id INTEGER PRIMARY KEY AUTOINCREMENT,

                hn TEXT,

                dispense_date TEXT,

                patient_json TEXT,

                status TEXT DEFAULT 'pending',

                created_at TEXT,

                resolved_at TEXT

            )
        """)


        # ====================================================
        # MEDICATION HISTORY
        # ====================================================

        conn.execute("""
            CREATE TABLE IF NOT EXISTS medication_history (

                id INTEGER PRIMARY KEY AUTOINCREMENT,

                queue_id INTEGER,

                hn TEXT,

                patient_name TEXT,

                age TEXT,

                dispense_date TEXT,

                appointment_date TEXT,

                decision TEXT,

                decision_text TEXT,

                original_drugs TEXT,

                substitute_drug TEXT,

                interaction_count INTEGER DEFAULT 0,

                dispensed_at TEXT

            )
        """)


        # ====================================================
        # MIGRATION QUEUE
        # ====================================================

        queue_columns = [

            ("hn", "TEXT"),

            ("dispense_date", "TEXT"),

            ("appointment_date", "TEXT"),

            ("patient_json", "TEXT"),

            ("status", "TEXT DEFAULT 'pending'"),

            ("created_at", "TEXT"),

            ("resolved_at", "TEXT")

        ]


        for column_name, column_type in queue_columns:

            add_column_if_missing(
                conn,
                "prescription_queue",
                column_name,
                column_type
            )


        # ====================================================
        # MIGRATION HISTORY
        # ====================================================

        history_columns = [

            ("queue_id", "INTEGER"),

            ("hn", "TEXT"),

            ("patient_name", "TEXT"),

            ("age", "TEXT"),

            ("dispense_date", "TEXT"),

            ("appointment_date", "TEXT"),

            ("decision", "TEXT"),

            ("decision_text", "TEXT"),

            ("original_drugs", "TEXT"),

            ("substitute_drug", "TEXT"),

            ("interaction_count", "INTEGER DEFAULT 0"),

            ("dispensed_at", "TEXT")

        ]


        for column_name, column_type in history_columns:

            add_column_if_missing(
                conn,
                "medication_history",
                column_name,
                column_type
            )


        # ====================================================
        # DEFAULT DATA FIX
        # ====================================================

        conn.execute("""
            UPDATE prescription_queue
            SET status = 'pending'
            WHERE status IS NULL
        """)


        conn.execute("""
            UPDATE medication_history
            SET interaction_count = 0
            WHERE interaction_count IS NULL
        """)


        # ====================================================
        # MERGE DUPLICATE PENDING QUEUE BY HN
        # ====================================================
        # 1 HN = 1 รายการบนหน้าผลการตรวจ
        # ถ้ามีหลายแถว/หลายรายการยาใน HN เดียวกัน ให้รวมยาเข้าด้วยกัน
        # และเก็บ dispense_date / appointment_date ของ HN เดียวกันไว้ในรายการเดียว

        duplicate_rows = conn.execute("""
            SELECT hn
            FROM prescription_queue
            WHERE hn IS NOT NULL AND status = 'pending'
            GROUP BY hn
            HAVING COUNT(*) > 1
        """).fetchall()

        for row in duplicate_rows:
            hn_value = str(row["hn"]).strip()
            queue_rows = conn.execute("""
                SELECT *
                FROM prescription_queue
                WHERE hn = ? AND status = 'pending'
                ORDER BY id ASC
            """, (hn_value,)).fetchall()

            if not queue_rows:
                continue

            keep = queue_rows[0]
            try:
                merged = json.loads(keep["patient_json"] or "{}")
            except Exception:
                merged = {}

            merged_medicines = list(merged.get("medicines", []) or [])
            dispense = str(keep["dispense_date"] or merged.get("dispense_date", "")).strip()
            appointment = str(keep["appointment_date"] or merged.get("appointment_date", "")).strip()

            for extra in queue_rows[1:]:
                try:
                    extra_patient = json.loads(extra["patient_json"] or "{}")
                except Exception:
                    extra_patient = {}
                merged_medicines.extend(extra_patient.get("medicines", []) or [])
                if not dispense:
                    dispense = str(extra["dispense_date"] or extra_patient.get("dispense_date", "")).strip()
                if not appointment:
                    appointment = str(extra["appointment_date"] or extra_patient.get("appointment_date", "")).strip()

                conn.execute(
                    "DELETE FROM prescription_queue WHERE id = ?",
                    (extra["id"],)
                )

            merged["hn"] = hn_value
            merged["dispense_date"] = dispense
            merged["appointment_date"] = appointment
            merged["dispense_date_raw"] = dispense
            merged["appointment_date_raw"] = appointment
            merged["required_days"] = calculate_days_between(dispense, appointment)
            merged["medicines"] = merged_medicines

            conn.execute("""
                UPDATE prescription_queue
                SET dispense_date = ?, appointment_date = ?, patient_json = ?
                WHERE id = ?
            """, (
                dispense,
                appointment,
                json.dumps(merged, ensure_ascii=False, default=str),
                keep["id"]
            ))

        # ลบ unique index เดิมที่บังคับ HN+วันที่ แล้วใช้ HN เป็นตัวหลัก
        conn.execute("DROP INDEX IF EXISTS uq_queue_hn_date")


        # ====================================================
        # DUPLICATE HISTORY CLEANUP
        # ====================================================

        duplicate_history = conn.execute("""
            SELECT
                queue_id,
                MAX(id) AS keep_id
            FROM medication_history
            WHERE queue_id IS NOT NULL
            GROUP BY queue_id
            HAVING COUNT(*) > 1
        """).fetchall()


        for row in duplicate_history:

            conn.execute("""
                DELETE FROM medication_history
                WHERE queue_id = ?
                AND id != ?
            """, (
                row["queue_id"],
                row["keep_id"]
            ))


        # ====================================================
        # INDEX
        # ====================================================

        conn.execute("""
            CREATE INDEX IF NOT EXISTS idx_queue_status
            ON prescription_queue(status)
        """)


        conn.execute("""
            CREATE INDEX IF NOT EXISTS idx_queue_hn
            ON prescription_queue(hn)
        """)


        conn.execute("""
            CREATE INDEX IF NOT EXISTS idx_history_queue
            ON medication_history(queue_id)
        """)


        # ====================================================
        # UNIQUE QUEUE
        # ====================================================

        conn.execute("""
            CREATE UNIQUE INDEX IF NOT EXISTS
            uq_queue_hn_pending
            ON prescription_queue(hn)
            WHERE status = 'pending'
        """)


        # ====================================================
        # UNIQUE HISTORY
        # ====================================================

        conn.execute("""
            CREATE UNIQUE INDEX IF NOT EXISTS
            uq_history_queue
            ON medication_history(queue_id)
        """)


        conn.commit()


        print("")
        print("=" * 70)
        print("DATABASE INITIALIZED SUCCESSFULLY")
        print("Database:", DATABASE_FILE)
        print("=" * 70)


    except Exception as e:

        conn.rollback()

        print("")
        print("=" * 70)
        print("ERROR INITIALIZING DATABASE")
        print(repr(e))
        print("=" * 70)

        raise


    finally:

        conn.close()


# ============================================================
# INITIALIZE DATABASE
# ============================================================

init_database()


# ============================================================
# TEST SUBSTITUTE DRUGS
# ============================================================

TEST_SUBSTITUTE_DRUGS = [

    {
        "name": "Test Substitute Drug A",
        "strength": "Test strength",
        "note": "ข้อมูลสำหรับทดสอบระบบเท่านั้น"
    },

    {
        "name": "Test Substitute Drug B",
        "strength": "Test strength",
        "note": "ข้อมูลสำหรับทดสอบระบบเท่านั้น"
    },

    {
        "name": "Test Substitute Drug C",
        "strength": "Test strength",
        "note": "ข้อมูลสำหรับทดสอบระบบเท่านั้น"
    }

]


# ============================================================
# NORMALIZE DRUG NAME
# ============================================================

def normalize_drug_name(name):

    if name is None:

        return ""


    name = str(name).strip().lower()

    name = name.replace(" ", "")
    name = name.replace("-", "")
    name = name.replace("_", "")


    return name


# ============================================================
# LOAD INTERACTIONS
# ============================================================

def load_interactions():

    interactions = []


    if not os.path.exists(
        INTERACTION_FILE
    ):

        return interactions


    try:

        wb = load_workbook(
            INTERACTION_FILE,
            data_only=True
        )


        ws = wb.active


        headers = [
            cell.value
            for cell in ws[1]
        ]


        for row in ws.iter_rows(
            min_row=2,
            values_only=True
        ):

            if not any(
                value is not None
                for value in row
            ):

                continue


            item = {}


            for i, header in enumerate(headers):

                if (
                    header is not None
                    and
                    i < len(row)
                ):

                    item[header] = row[i]


            interactions.append(item)


        wb.close()


    except Exception as e:

        print(
            "ERROR loading interaction file:",
            repr(e)
        )


    return interactions


# ============================================================
# CHECK DRUG INTERACTIONS
# ============================================================

def check_drug_interactions(
    drug_names
):

    interactions = load_interactions()
    results = []

    normalized_names = [
        normalize_drug_name(name)
        for name in drug_names
        if name
    ]

    for interaction in interactions:
        drug1 = interaction.get("Drug_1", "")
        drug2 = interaction.get("Drug_2", "")

        drug1_norm = normalize_drug_name(drug1)
        drug2_norm = normalize_drug_name(drug2)

        if (
            drug1_norm in normalized_names
            and drug2_norm in normalized_names
        ) or (
            drug2_norm in normalized_names
            and drug1_norm in normalized_names
        ):
            results.append(interaction)

    return results


# ============================================================
# CHECK DRUG INTERACTIONS BY ACTUAL MEDICATION DATES
# ============================================================

def check_patient_drug_interactions(patient):
    """
    ตรวจ Drug Interaction เฉพาะยาที่มีช่วงการใช้ยาทับซ้อนกันจริง
    ของ HN เดียวกัน

    ลำดับ:
        1. ใช้ dispense_date ของ HN เป็นวันเริ่มยา
        2. คำนวณวันใช้ยาจาก quantity / times_per_day
        3. หาวันสิ้นสุดของยาแต่ละตัว
        4. ตรวจว่าช่วงวันของยาคู่ใดทับซ้อนกันหรือไม่
        5. ถ้าทับซ้อน จึงนำคู่นั้นไปตรวจใน drug_interactions.xlsx
    """
    medicines = patient.get("medicines", []) or []
    interactions = load_interactions()

    if len(medicines) < 2:
        return []

    default_start = parse_date_value(
        patient.get("dispense_date", "")
    )

    if default_start is None:
        return []

    # สร้างช่วงวันใช้ยาของแต่ละรายการ
    medication_ranges = []

    for medicine in medicines:
        name = str(medicine.get("name", "") or "").strip()
        if not name:
            continue

        start = parse_date_value(
            medicine.get("start_date", "")
        ) or default_start

        try:
            quantity = float(medicine.get("quantity", 0) or 0)
        except Exception:
            quantity = 0

        try:
            times = float(medicine.get("times_per_day", 0) or 0)
        except Exception:
            times = 0

        # ถ้าข้อมูลยังไม่มี days_supply ให้คำนวณใหม่จากจำนวนยา/ครั้งต่อวัน
        try:
            days_supply = float(medicine.get("days_supply", 0) or 0)
        except Exception:
            days_supply = 0

        if days_supply <= 0 and quantity > 0 and times > 0:
            days_supply = quantity / times

        if days_supply <= 0:
            continue

        import math
        duration_days = max(1, math.ceil(days_supply))
        end = start + __import__("datetime").timedelta(days=duration_days - 1)

        medication_ranges.append({
            "name": name,
            "start": start,
            "end": end
        })

    if len(medication_ranges) < 2:
        return []

    # ตรวจคู่ยาเฉพาะคู่ที่ช่วงวันทับซ้อนกัน
    active_pairs = set()
    for i in range(len(medication_ranges)):
        for j in range(i + 1, len(medication_ranges)):
            a = medication_ranges[i]
            b = medication_ranges[j]

            overlap_start = max(a["start"], b["start"])
            overlap_end = min(a["end"], b["end"])

            if overlap_start <= overlap_end:
                pair = frozenset((
                    normalize_drug_name(a["name"]),
                    normalize_drug_name(b["name"])
                ))
                active_pairs.add(pair)

    if not active_pairs:
        return []

    results = []
    seen = set()

    for interaction in interactions:
        drug1 = interaction.get("Drug_1", "")
        drug2 = interaction.get("Drug_2", "")
        drug1_norm = normalize_drug_name(drug1)
        drug2_norm = normalize_drug_name(drug2)

        pair = frozenset((drug1_norm, drug2_norm))
        if pair in active_pairs and pair not in seen:
            results.append(interaction)
            seen.add(pair)

    return results


# ============================================================
# FIND COLUMN
# ============================================================

def find_column(
    headers,
    possible_names
):

    normalized_headers = {}


    for header in headers:

        if header is None:

            continue


        key = str(
            header
        ).strip().lower()


        normalized_headers[key] = header


    for name in possible_names:

        key = str(
            name
        ).strip().lower()


        if key in normalized_headers:

            return normalized_headers[key]


    return None


# ============================================================
# DATE HELPERS
# ============================================================

def parse_date_value(value):
    """แปลงวันที่จาก Excel/ข้อความให้เป็น datetime อย่างปลอดภัย"""
    if value is None or value == "":
        return None

    if isinstance(value, datetime):
        return value

    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())

    # Excel serial date
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            number = float(value)
            if number > 20000:
                return from_excel(number)
        except Exception:
            pass

    text = str(value).strip()
    if not text:
        return None

    # รองรับวันที่ที่มาจาก Excel และข้อความหลายรูปแบบ
    formats = [
        "%d/%m/%Y", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M",
        "%d-%m-%Y", "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M",
        "%Y/%m/%d", "%Y/%m/%d %H:%M:%S", "%Y/%m/%d %H:%M",
        "%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M",
        "%m/%d/%Y", "%m/%d/%Y %H:%M:%S", "%m/%d/%Y %H:%M",
    ]

    for fmt in formats:
        try:
            return datetime.strptime(text, fmt)
        except Exception:
            pass

    return None


def format_date(value):
    parsed = parse_date_value(value)
    if parsed is None:
        return "" if value is None else str(value).strip()
    return parsed.strftime("%d/%m/%Y")


# ============================================================
# DATE SORT KEY
# ============================================================

def date_sort_key(value):
    parsed = parse_date_value(value)
    return parsed if parsed is not None else datetime.min


# ============================================================
# PARSE FREQUENCY / TIMES PER DAY
# ============================================================

def parse_times_per_day(value, default=1):
    """แปลงครั้งต่อวันจาก Excel ให้เป็นตัวเลข แม้จะเขียนเป็น 1x, 1-0-1 หรือ วันละ 2 ครั้ง"""
    if value is None or value == "":
        return default

    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            n = float(value)
            return int(n) if n.is_integer() else n
        except Exception:
            return default

    text = str(value).strip().lower()
    if not text:
        return default

    # ตัวเลขตรง ๆ เช่น 1, 2, 3
    try:
        n = float(text)
        return int(n) if n.is_integer() else n
    except Exception:
        pass

    # รูปแบบ 1-0-1 / 1-1-1 / 0-1-1 ให้รวมตัวเลขที่คั่นด้วย - หรือ /
    import re
    nums = re.findall(r"\d+(?:\.\d+)?", text)
    if ("-" in text or "/" in text or "x" in text) and nums:
        # ถ้าเป็น 1x2 ให้ใช้เลขหลัง x; ถ้าเป็น 1-0-1 ให้รวมทุกช่วง
        if "x" in text and len(nums) >= 2:
            try:
                return float(nums[-1])
            except Exception:
                pass
        try:
            total = sum(float(x) for x in nums)
            return int(total) if total.is_integer() else total
        except Exception:
            pass

    # ข้อความ เช่น "วันละ 2 ครั้ง"
    if nums:
        try:
            n = float(nums[-1])
            return int(n) if n.is_integer() else n
        except Exception:
            pass

    return default


# ============================================================
# CALCULATE DAYS
# ============================================================

def calculate_days_between(start_date, end_date):
    start = parse_date_value(start_date)
    end = parse_date_value(end_date)

    if start is None or end is None:
        return 0

    days = (end.date() - start.date()).days

    # วันจ่ายและวันนัดวันเดียวกัน = ต้องใช้ยา 1 วัน ไม่ใช่ 0 วัน
    if days == 0:
        return 1

    return days


# ============================================================
# CHECK DAYS SUPPLY
# ============================================================

def calculate_days_check(patient):
    """
    ตรวจสอบจำนวนยาให้สัมพันธ์กับช่วงวันที่จ่ายยา -> วันนัด

    หลักการ:
        required_days = วันนัด - วันที่จ่ายยา
        expected_quantity = required_days x ครั้ง/วัน
        missing_quantity = expected_quantity - จำนวนที่สั่ง  (กรณีขาด)
        excess_quantity = จำนวนที่สั่ง - expected_quantity  (กรณีเกิน)

    ระบบจะแจ้งเป็น "ขาดกี่เม็ด / เกินกี่เม็ด" ไม่ใช่เพียงจำนวนวัน
    """

    # คำนวณใหม่จากวันที่ของ HN นี้ทุกครั้ง ไม่ใช้ค่าค้างจากแถวอื่น
    required_days = calculate_days_between(
        patient.get("dispense_date", ""),
        patient.get("appointment_date", "")
    )
    patient["required_days"] = required_days

    try:
        required_days = float(required_days or 0)
    except Exception:
        required_days = 0

    results = []
    has_problem = False

    for medicine in patient.get("medicines", []):

        try:
            quantity = float(
                medicine.get("quantity", 0) or 0
            )
        except Exception:
            quantity = 0

        try:
            times = float(
                medicine.get("times_per_day", 1) or 1
            )
        except Exception:
            times = 0

        missing_quantity = 0
        excess_quantity = 0

        # ----------------------------------------------------
        # INVALID DATE / FREQUENCY
        # ----------------------------------------------------
        if required_days <= 0 or times <= 0:

            expected = 0
            decision = "consult"
            status = "คำนวณไม่ได้ — กรุณาตรวจสอบวันที่จ่าย/วันนัด/ครั้งต่อวัน"
            has_problem = True

        else:

            # จำนวนเม็ดที่ควรจ่ายเพื่อให้ถึงวันนัด
            expected = required_days * times

            if quantity < expected:

                missing_quantity = expected - quantity
                excess_quantity = 0
                decision = "increase"
                status = "ขาด {} เม็ด".format(
                    format_quantity(missing_quantity)
                )
                has_problem = True

            elif quantity > expected:

                missing_quantity = 0
                excess_quantity = quantity - expected
                decision = "decrease"
                status = "เกิน {} เม็ด".format(
                    format_quantity(excess_quantity)
                )
                has_problem = True

            else:

                decision = "correct"
                status = "พอดี"

        expected_display = (
            int(expected)
            if isinstance(expected, (int, float))
            and float(expected).is_integer()
            else expected
        )

        quantity_display = (
            int(quantity)
            if float(quantity).is_integer()
            else quantity
        )

        times_display = (
            int(times)
            if float(times).is_integer()
            else times
        )

        missing_display = (
            int(missing_quantity)
            if float(missing_quantity).is_integer()
            else missing_quantity
        )

        excess_display = (
            int(excess_quantity)
            if float(excess_quantity).is_integer()
            else excess_quantity
        )

        medicine["expected_quantity"] = expected_display
        medicine["missing_quantity"] = missing_display
        medicine["excess_quantity"] = excess_display
        medicine["days_match"] = decision == "correct"
        medicine["days_check_status"] = status
        medicine["required_days"] = required_days
        medicine["times_per_day"] = times_display

        results.append({
            "name": medicine.get("name", ""),
            "strength": medicine.get("strength", ""),
            "quantity": quantity_display,
            "times_per_day": times_display,
            "required_days": required_days,
            "expected_quantity": expected_display,
            "missing_quantity": missing_display,
            "excess_quantity": excess_display,
            "status": status,
            "decision": decision
        })

    patient["days_check_results"] = results

    # --------------------------------------------------------
    # ถ้าแพทย์/ผู้มีอำนาจเคยพิจารณาแล้ว ให้ถือว่าผ่าน Consult
    # --------------------------------------------------------
    manual_decision = patient.get(
        "_manual_days_decision"
    )

    if manual_decision in {
        "increase",
        "decrease",
        "correct",
        "override_doctor"
    }:
        patient["days_check_required"] = False
    else:
        patient["days_check_required"] = has_problem

    # ข้อความสรุปสำหรับหน้าเว็บ / log / PDF
    problem_items = [
        item for item in results
        if item.get("decision") in {
            "increase", "decrease", "consult"
        }
    ]

    patient["days_supply_problem_count"] = len(problem_items)
    patient["days_supply_has_problem"] = len(problem_items) > 0

    missing_total = 0
    excess_total = 0

    for item in problem_items:
        try:
            missing_total += float(
                item.get("missing_quantity", 0) or 0
            )
        except Exception:
            pass
        try:
            excess_total += float(
                item.get("excess_quantity", 0) or 0
            )
        except Exception:
            pass

    patient["days_supply_missing_total"] = (
        int(missing_total)
        if float(missing_total).is_integer()
        else missing_total
    )
    patient["days_supply_excess_total"] = (
        int(excess_total)
        if float(excess_total).is_integer()
        else excess_total
    )

    return patient


# ============================================================
# FORMAT QUANTITY FOR DISPLAY
# ============================================================

def format_quantity(value):
    try:
        number = float(value)
        if number.is_integer():
            return str(int(number))
        return "{:.2f}".format(number).rstrip("0").rstrip(".")
    except Exception:
        return str(value)


# ============================================================
# READ PRESCRIPTION EXCEL
# ============================================================

def read_prescription_excel(
    file_path
):

    patients = {}


    wb = load_workbook(
        file_path,
        data_only=True
    )


    ws = wb.active


    rows = list(
        ws.iter_rows(
            values_only=True
        )
    )


    if not rows:

        wb.close()

        return []


    headers = list(
        rows[0]
    )


    # ========================================================
    # FIND COLUMNS
    # ========================================================

    hn_col = find_column(
        headers,
        [
            "HN",
            "hn",
            "Hospital Number",
            "hospital_number"
        ]
    )


    name_col = find_column(
        headers,
        [
            "Name",
            "name",
            "Patient Name",
            "patient_name",
            "ชื่อ"
        ]
    )


    age_col = find_column(
        headers,
        [
            "Age",
            "age",
            "อายุ"
        ]
    )


    drug_col = find_column(
        headers,
        [
            "Drug",
            "drug",
            "Drug Name",
            "drug_name",
            "Medicine",
            "medicine",
            "ชื่อยา"
        ]
    )


    strength_col = find_column(
        headers,
        [
            "Strength",
            "strength",
            "Dose",
            "dose",
            "ขนาดยา"
        ]
    )


    times_col = find_column(
        headers,
        [
            "Times Per Day",
            "times_per_day",
            "times",
            "Frequency",
            "frequency",
            "ครั้งต่อวัน",
            "ครั้ง/วัน"
        ]
    )


    quantity_col = find_column(
        headers,
        [
            "Quantity",
            "quantity",
            "Qty",
            "qty",
            "จำนวน"
        ]
    )


    dispense_col = find_column(
        headers,
        [
            "Dispense Date",
            "dispense_date",
            "Date",
            "date",
            "วันที่จ่ายยา"
        ]
    )


    appointment_col = find_column(
        headers,
        [
            "Appointment Date",
            "appointment_date",
            "Appointment",
            "appointment",
            "วันนัด"
        ]
    )


    if hn_col is None:

        wb.close()

        raise ValueError(
            "ไม่พบ column HN ในไฟล์ Excel"
        )


    # ========================================================
    # READ DATA
    # ========================================================

    for row in rows[1:]:

        row_dict = {}


        for i, header in enumerate(headers):

            if header is None:

                continue


            if i < len(row):

                row_dict[header] = row[i]


        hn = row_dict.get(
            hn_col
        )


        if hn is None:

            continue


        if (
            isinstance(hn, float)
            and
            hn.is_integer()
        ):

            hn = str(
                int(hn)
            )

        else:

            hn = str(
                hn
            ).strip()


        if not hn:

            continue


        # ====================================================
        # DATE ของแถวปัจจุบัน -- ต้องอ่านใหม่ทุกแถว
        # เพื่อไม่ให้ HN หนึ่งใช้วันที่จากแถวก่อนหน้า/คนอื่น
        # ====================================================
        dispense_date_raw = (
            row_dict.get(dispense_col, "") if dispense_col else ""
        )
        appointment_date_raw = (
            row_dict.get(appointment_col, "") if appointment_col else ""
        )

        # ====================================================
        # CREATE PATIENT
        # ====================================================

        if hn not in patients:

            patient_name = ""


            if name_col:

                value = row_dict.get(
                    name_col,
                    ""
                )


                if value is not None:

                    patient_name = str(
                        value
                    )


            age = ""


            if age_col:

                value = row_dict.get(
                    age_col,
                    ""
                )


                if value is not None:

                    age = str(
                        value
                    )


            required_days = calculate_days_between(
                dispense_date_raw,
                appointment_date_raw
            )


            patients[hn] = {

                "hn": hn,

                "name": patient_name,

                "age": age,

                "dispense_date_raw":
                    dispense_date_raw,

                "appointment_date_raw":
                    appointment_date_raw,

                "dispense_date":
                    format_date(
                        dispense_date_raw
                    ),

                "appointment_date":
                    format_date(
                        appointment_date_raw
                    ),

                "required_days":
                    required_days,

                "medicines": [],

                "interaction_results": []

            }

        else:
            # HN เดิม: ใช้ข้อมูลวันของ HN เดียวกันจากแถวที่มีข้อมูลจริง
            patient = patients[hn]
            if not patient.get("dispense_date") and dispense_date_raw:
                formatted = format_date(dispense_date_raw)
                if formatted:
                    patient["dispense_date_raw"] = dispense_date_raw
                    patient["dispense_date"] = formatted
            if not patient.get("appointment_date") and appointment_date_raw:
                formatted = format_date(appointment_date_raw)
                if formatted:
                    patient["appointment_date_raw"] = appointment_date_raw
                    patient["appointment_date"] = formatted
            patient["required_days"] = calculate_days_between(
                patient.get("dispense_date", ""),
                patient.get("appointment_date", "")
            )


        # ====================================================
        # DRUG
        # ====================================================

        drug_name = ""


        if drug_col:

            value = row_dict.get(
                drug_col
            )


            if value is not None:

                drug_name = str(
                    value
                ).strip()


        if not drug_name:

            continue


        # ====================================================
        # STRENGTH
        # ====================================================

        strength = ""


        if strength_col:

            value = row_dict.get(
                strength_col
            )


            if value is not None:

                strength = str(
                    value
                )


        # ====================================================
        # TIMES
        # ====================================================

        times = parse_times_per_day(
            row_dict.get(times_col, 1) if times_col else 1,
            default=1
        )


        # ====================================================
        # QUANTITY
        # ====================================================

        quantity = 0


        if quantity_col:

            value = row_dict.get(
                quantity_col
            )


            if value is not None:

                try:

                    quantity = float(
                        value
                    )


                    if quantity.is_integer():

                        quantity = int(
                            quantity
                        )


                except Exception:

                    quantity = 0


        # ====================================================
        # DAYS SUPPLY
        # ====================================================

        days_supply = 0


        try:

            numeric_quantity = float(
                quantity
            )


            numeric_times = float(
                times
            )


            if numeric_times > 0:

                days_supply = (
                    numeric_quantity
                    /
                    numeric_times
                )


        except Exception:

            days_supply = 0


        # ====================================================
        # INITIAL STATUS
        # ====================================================

        required_days = patients[hn][
            "required_days"
        ]


        if days_supply >= required_days:

            status = "เพียงพอ"

        else:

            status = "ไม่เพียงพอ"


        patients[hn]["medicines"].append({

            "name": drug_name,

            "strength": strength,

            "times_per_day": times,

            "quantity": quantity,

            "days_supply": days_supply,

            "status": status

        })


    wb.close()


    # ========================================================
    # INTERACTION
    # ========================================================

    patient_results = []


    for hn, patient in patients.items():

        drug_names = [

            medicine["name"]

            for medicine
            in patient["medicines"]

        ]


        # ====================================================
        # CALCULATE DAYS FIRST
        # ====================================================
        # คำนวณวันที่จ่าย -> วันนัดของ HN นี้ก่อนทุกครั้ง
        # เพื่อให้ตรวจสอบจำนวนยา/ช่วงวันของ HN นี้เสร็จก่อน
        patient = calculate_days_check(patient)

        # ====================================================
        # INTERACTION CHECK AFTER DAYS CHECK
        # ====================================================
        patient["interaction_results"] = check_patient_drug_interactions(
            patient
        )


        patient_results.append(
            patient
        )


    # ========================================================
    # SORT OLD → NEW
    # ========================================================

    patient_results.sort(

        key=lambda p:
        date_sort_key(
            p.get(
                "dispense_date",
                ""
            )
        )

    )


    return patient_results


# ============================================================
# LOAD STOCK
# ============================================================

def load_stock():

    stock_data = []


    if not os.path.exists(
        STOCK_FILE
    ):

        return stock_data


    try:

        wb = load_workbook(
            STOCK_FILE,
            data_only=True
        )


        ws = wb.active


        headers = [

            cell.value

            for cell in ws[1]

        ]


        for row in ws.iter_rows(
            min_row=2,
            values_only=True
        ):

            if not any(
                value is not None
                for value in row
            ):

                continue


            item = {}


            for i, header in enumerate(headers):

                if (
                    header is not None
                    and
                    i < len(row)
                ):

                    item[header] = row[i]


            stock_data.append(
                item
            )


        wb.close()


    except Exception as e:

        print(
            "ERROR LOAD STOCK:",
            repr(e)
        )


    return stock_data


# ============================================================
# FIND STOCK ITEM
# ============================================================

def find_stock_item(
    drug_name,
    stock_data
):

    target = normalize_drug_name(
        drug_name
    )


    stock_name_keys = [

        "Generic_Name",

        "Generic Name",

        "generic_name",

        "ชื่อยา",

        "ชื่อยา Generic",

        "Drug",

        "Drug Name",

        "Drug_Name",

        "Medicine",

        "Medicine Name"

    ]


    for item in stock_data:

        for key in stock_name_keys:

            if key not in item:

                continue


            stock_name = normalize_drug_name(
                item.get(key)
            )


            if stock_name == target:

                return item


    return None


# ============================================================
# FIND STOCK ROW
# ============================================================

def find_stock_row(
    ws,
    drug_name
):

    headers = [
        cell.value
        for cell in ws[1]
    ]


    name_column = None


    possible_name_columns = [

        "Generic_Name",

        "Generic Name",

        "generic_name",

        "ชื่อยา",

        "ชื่อยา Generic",

        "Drug",

        "Drug Name",

        "Drug_Name",

        "Medicine",

        "Medicine Name"

    ]


    # ========================================================
    # FIND DRUG NAME COLUMN
    # ========================================================

    for index, header in enumerate(
        headers,
        start=1
    ):

        if header is None:

            continue


        header_text = str(
            header
        ).strip().lower()


        for possible in possible_name_columns:

            possible_text = str(
                possible
            ).strip().lower()


            if header_text == possible_text:

                name_column = index

                break


        if name_column is not None:

            break


    if name_column is None:

        return None, None


    # ========================================================
    # FIND DRUG ROW
    # ========================================================

    target = normalize_drug_name(
        drug_name
    )


    for row_number in range(
        2,
        ws.max_row + 1
    ):

        value = ws.cell(
            row=row_number,
            column=name_column
        ).value


        stock_name = normalize_drug_name(
            value
        )


        if stock_name == target:

            return row_number, name_column


    return None, name_column


# ============================================================
# FIND STOCK QUANTITY COLUMN
# ============================================================

def find_stock_quantity_column(
    ws
):

    headers = [

        cell.value

        for cell in ws[1]

    ]


    possible_quantity_columns = [

        "Quantity",

        "quantity",

        "Qty",

        "qty",

        "จำนวน",

        "จำนวนคงเหลือ",

        "คงเหลือ",

        "Stock",

        "Stock Quantity",

        "stock_quantity"

    ]


    for index, header in enumerate(
        headers,
        start=1
    ):

        if header is None:

            continue


        header_text = str(
            header
        ).strip().lower()


        for possible in possible_quantity_columns:

            possible_text = str(
                possible
            ).strip().lower()


            if header_text == possible_text:

                return index


    return None


# ============================================================
# FIND OR CREATE HEADER
# ============================================================

def find_or_create_header(
    ws,
    header_name
):

    for column in range(
        1,
        ws.max_column + 1
    ):

        value = ws.cell(
            row=1,
            column=column
        ).value


        if value is None:

            continue


        if str(
            value
        ).strip().lower() == str(
            header_name
        ).strip().lower():

            return column


    new_column = ws.max_column + 1


    ws.cell(
        row=1,
        column=new_column,
        value=header_name
    )


    return new_column


# ============================================================
# DEDUCT STOCK FROM EXCEL
# ============================================================

def deduct_stock_from_excel(
    patient
):

    """
    หักจำนวนยาใน stock.xlsx

    ก่อนหักจะตรวจสอบยาทุกตัวก่อน
    ถ้ามียาตัวใดไม่พบหรือไม่พอ
    จะไม่หัก Stock ตัวใดเลย
    """


    # ========================================================
    # CHECK STOCK FILE
    # ========================================================

    if not os.path.exists(
        STOCK_FILE
    ):

        print(
            "ERROR: ไม่พบไฟล์ stock.xlsx"
        )

        return False


    wb = None


    try:

        # ====================================================
        # OPEN STOCK
        # ====================================================

        wb = load_workbook(
            STOCK_FILE
        )


        ws = wb.active


        # ====================================================
        # FIND QUANTITY COLUMN
        # ====================================================

        quantity_column = (
            find_stock_quantity_column(
                ws
            )
        )


        if quantity_column is None:

            raise ValueError(
                "ไม่พบ Column Quantity ใน stock.xlsx"
            )


        # ====================================================
        # CREATE LOG COLUMNS
        # ====================================================

        last_dispensed_at_column = (
            find_or_create_header(

                ws,

                "Last_Dispensed_At"

            )
        )


        last_dispensed_quantity_column = (
            find_or_create_header(

                ws,

                "Last_Dispensed_Quantity"

            )
        )


        # ====================================================
        # CURRENT DISPENSE TIME
        # ====================================================

        dispense_time = datetime.now().strftime(
            "%Y-%m-%d %H:%M:%S"
        )


        # ====================================================
        # GET MEDICINES
        # ====================================================

        medicines = patient.get(
            "medicines",
            []
        )


        if not medicines:

            raise ValueError(
                "ไม่พบรายการยาในใบสั่งยา"
            )


        # ====================================================
        # AGGREGATE SAME DRUG
        #
        # หากใบสั่งยามียาตัวเดียวกันหลายรายการ
        # จะรวมจำนวนก่อนหัก
        # ====================================================

        required_by_drug = {}


        display_names = {}


        for medicine in medicines:

            drug_name = str(
                medicine.get(
                    "name",
                    ""
                )
            ).strip()


            if not drug_name:

                raise ValueError(
                    "พบรายการยาที่ไม่มีชื่อยา"
                )


            try:

                required_quantity = float(
                    medicine.get(
                        "quantity",
                        0
                    ) or 0
                )


            except Exception:

                raise ValueError(
                    "จำนวนยาไม่ถูกต้อง: "
                    + drug_name
                )


            if required_quantity <= 0:

                raise ValueError(
                    "จำนวนยาต้องมากกว่า 0: "
                    + drug_name
                )


            normalized = normalize_drug_name(
                drug_name
            )


            if normalized not in required_by_drug:

                required_by_drug[normalized] = 0


            required_by_drug[normalized] += (
                required_quantity
            )


            display_names[normalized] = (
                drug_name
            )


        # ====================================================
        # VALIDATE ALL STOCK FIRST
        # ====================================================

        deductions = []


        for normalized_name, required_quantity in (
            required_by_drug.items()
        ):

            drug_name = display_names[
                normalized_name
            ]


            row_number, name_column = (
                find_stock_row(

                    ws,

                    drug_name

                )
            )


            if row_number is None:

                raise ValueError(
                    "ไม่พบยาใน stock.xlsx: "
                    + drug_name
                )


            # ------------------------------------------------
            # CURRENT STOCK
            # ------------------------------------------------

            current_value = ws.cell(
                row=row_number,
                column=quantity_column
            ).value


            try:

                current_quantity = float(
                    current_value or 0
                )


            except Exception:

                raise ValueError(
                    "จำนวน Stock ไม่ถูกต้อง: "
                    + drug_name
                )


            # ------------------------------------------------
            # CHECK SUFFICIENT
            # ------------------------------------------------

            if current_quantity < required_quantity:

                raise ValueError(

                    "Stock ไม่เพียงพอ: {} "
                    "(เหลือ {}, ต้องใช้ {})"

                    .format(

                        drug_name,

                        current_quantity,

                        required_quantity

                    )

                )


            # ------------------------------------------------
            # NEW STOCK
            # ------------------------------------------------

            new_quantity = (
                current_quantity
                -
                required_quantity
            )


            deductions.append({

                "drug_name":
                    drug_name,

                "row":
                    row_number,

                "old_quantity":
                    current_quantity,

                "required_quantity":
                    required_quantity,

                "new_quantity":
                    new_quantity

            })


        # ====================================================
        # APPLY DEDUCTIONS
        # ====================================================

        for item in deductions:

            row_number = item[
                "row"
            ]


            new_quantity = item[
                "new_quantity"
            ]


            required_quantity = item[
                "required_quantity"
            ]


            # ------------------------------------------------
            # INTEGER FORMAT
            # ------------------------------------------------

            if float(
                new_quantity
            ).is_integer():

                new_quantity = int(
                    new_quantity
                )


            if float(
                required_quantity
            ).is_integer():

                required_quantity = int(
                    required_quantity
                )


            # ------------------------------------------------
            # UPDATE QUANTITY
            # ------------------------------------------------

            ws.cell(

                row=row_number,

                column=quantity_column,

                value=new_quantity

            )


            # ------------------------------------------------
            # UPDATE DISPENSE TIME
            # ------------------------------------------------

            ws.cell(

                row=row_number,

                column=last_dispensed_at_column,

                value=dispense_time

            )


            # ------------------------------------------------
            # UPDATE DISPENSE QUANTITY
            # ------------------------------------------------

            ws.cell(

                row=row_number,

                column=last_dispensed_quantity_column,

                value=required_quantity

            )


        # ====================================================
        # SAVE STOCK
        # ====================================================

        wb.save(
            STOCK_FILE
        )


        wb.close()

        wb = None


        # ====================================================
        # PRINT LOG
        # ====================================================

        print("")
        print("=" * 70)
        print("STOCK DEDUCTED SUCCESSFULLY")
        print(
            "Dispense Time:",
            dispense_time
        )


        for item in deductions:

            print(

                "{} | {} -> {} | จ่าย {}"

                .format(

                    item["drug_name"],

                    item["old_quantity"],

                    item["new_quantity"],

                    item["required_quantity"]

                )

            )


        print("=" * 70)


        return True


    except Exception as e:

        print("")
        print("=" * 70)
        print("ERROR DEDUCT STOCK")
        print(
            repr(e)
        )
        print("=" * 70)


        if wb is not None:

            try:

                wb.close()

            except Exception:

                pass


        return False


# ============================================================
# CHECK STOCK FOR ONE DRUG
# ============================================================

def check_stock_drug(
    drug_name,
    required_quantity=0
):

    stock_data = load_stock()


    item = find_stock_item(
        drug_name,
        stock_data
    )


    # ========================================================
    # NOT FOUND
    # ========================================================

    if item is None:

        return {

            "found": False,

            "matched_name": drug_name,

            "strength": "",

            "dosage_form": "",

            "quantity": 0,

            "required": required_quantity,

            "unit": "",

            "min_stock": 0,

            "status":
                "ไม่พบข้อมูลยาใน stock.xlsx"

        }


    # ========================================================
    # FIND QUANTITY
    # ========================================================

    quantity = 0


    quantity_keys = [

        "Quantity",

        "quantity",

        "Qty",

        "qty",

        "จำนวน",

        "จำนวนคงเหลือ",

        "คงเหลือ",

        "Stock",

        "Stock Quantity",

        "stock_quantity"

    ]


    for key in quantity_keys:

        if key in item:

            quantity = item.get(
                key
            )

            break


    try:

        current_quantity = float(
            quantity or 0
        )

    except Exception:

        current_quantity = 0


    try:

        required = float(
            required_quantity or 0
        )

    except Exception:

        required = 0


    # ========================================================
    # STOCK STATUS
    # ========================================================

    if current_quantity >= required:

        status = "มีเพียงพอ"

    else:

        status = "ยาใน Stock ไม่เพียงพอ"


    # ========================================================
    # MATCHED NAME
    # ========================================================

    matched_name = drug_name


    for key in [

        "Generic_Name",

        "Generic Name",

        "generic_name",

        "ชื่อยา",

        "ชื่อยา Generic",

        "Drug",

        "Drug Name",

        "Drug_Name",

        "Medicine",

        "Medicine Name"

    ]:

        if key in item:

            matched_name = item.get(
                key
            )

            break


    # ========================================================
    # STRENGTH
    # ========================================================

    strength = ""


    for key in [

        "Strength",

        "strength",

        "Dose",

        "dose",

        "ขนาดยา"

    ]:

        if key in item:

            strength = item.get(
                key
            )

            break


    # ========================================================
    # DOSAGE FORM
    # ========================================================

    dosage_form = ""


    for key in [

        "Dosage_Form",

        "Dosage Form",

        "dosage_form",

        "รูปแบบยา"

    ]:

        if key in item:

            dosage_form = item.get(
                key
            )

            break


    # ========================================================
    # UNIT
    # ========================================================

    unit = ""


    for key in [

        "Unit",

        "unit",

        "หน่วย"

    ]:

        if key in item:

            unit = item.get(
                key
            )

            break


    # ========================================================
    # MIN STOCK
    # ========================================================

    min_stock = 0


    for key in [

        "Min_Stock",

        "Min Stock",

        "min_stock",

        "Minimum Stock",

        "ขั้นต่ำ"

    ]:

        if key in item:

            min_stock = item.get(
                key
            )

            break


    return {

        "found": True,

        "matched_name":
            matched_name,

        "strength":
            strength,

        "dosage_form":
            dosage_form,

        "quantity":
            current_quantity,

        "required":
            required,

        "unit":
            unit,

        "min_stock":
            min_stock,

        "status":
            status

    }


# ============================================================
# CHECK PATIENT STOCK
# ============================================================

def check_patient_stock(
    patient
):

    results = []


    for medicine in patient.get(
        "medicines",
        []
    ):

        result = check_stock_drug(

            medicine.get(
                "name",
                ""
            ),

            medicine.get(
                "quantity",
                0
            )

        )


        results.append({

            "name":
                medicine.get(
                    "name",
                    ""
                ),

            "required":
                medicine.get(
                    "quantity",
                    0
                ),

            **result

        })


    return results


# ============================================================
# APPLY CURRENT STOCK STATUS TO MEDICINES
# ============================================================

def apply_stock_to_patient(
    patient
):

    stock_results = check_patient_stock(
        patient
    )


    patient["stock_results"] = stock_results


    for medicine in patient.get(
        "medicines",
        []
    ):

        medicine_name = normalize_drug_name(
            medicine.get(
                "name",
                ""
            )
        )


        stock = None


        for item in stock_results:

            stock_name = normalize_drug_name(
                item.get(
                    "name",
                    ""
                )
            )


            if stock_name == medicine_name:

                stock = item

                break


        # ====================================================
        # NOT FOUND
        # ====================================================

        if stock is None:

            medicine["stock_found"] = False

            medicine["stock_quantity"] = 0

            medicine["stock_required"] = medicine.get(
                "quantity",
                0
            )

            medicine["stock_status"] = (
                "ไม่พบข้อมูลยาใน stock.xlsx"
            )

            medicine["status"] = "ไม่เพียงพอ"

            continue


        # ====================================================
        # STOCK INFORMATION
        # ====================================================

        medicine["stock_found"] = stock.get(
            "found",
            False
        )


        medicine["stock_quantity"] = stock.get(
            "quantity",
            0
        )


        medicine["stock_required"] = stock.get(
            "required",
            medicine.get(
                "quantity",
                0
            )
        )


        medicine["stock_status"] = stock.get(
            "status",
            ""
        )


        medicine["stock_matched_name"] = stock.get(
            "matched_name",
            medicine.get(
                "name",
                ""
            )
        )


        medicine["stock_unit"] = stock.get(
            "unit",
            ""
        )


        # ====================================================
        # STOCK STATUS
        # ====================================================

        if (

            stock.get(
                "found",
                False
            )

            and

            stock.get(
                "quantity",
                0
            )

            >=

            stock.get(
                "required",
                0
            )

        ):

            medicine["status"] = "เพียงพอ"

        else:

            medicine["status"] = "ไม่เพียงพอ"


    return patient


# ============================================================
# SYNC PRESCRIPTION QUEUE FROM EXCEL
# ============================================================

def sync_prescription_queue_from_excel():
    """
    Sync Excel -> Queue โดยใช้ HN เป็นตัวหลัก

    HN เดียวกันใน Excel จะถูกรวมเป็นผู้ป่วย 1 ราย
    และส่ง dispense_date / appointment_date ของ HN นั้นไปพร้อมกันทุกครั้ง
    """
    if not os.path.exists(PRESCRIPTION_SOURCE_FILE):
        return False

    try:
        excel_patients = read_prescription_excel(PRESCRIPTION_SOURCE_FILE)
        source_map = {
            str(patient.get("hn", "")).strip(): patient
            for patient in excel_patients
            if str(patient.get("hn", "")).strip()
        }

        conn = get_db()
        try:
            pending_rows = conn.execute("""
                SELECT id, hn, dispense_date, appointment_date, patient_json
                FROM prescription_queue
                WHERE status = 'pending'
                ORDER BY id ASC
            """).fetchall()

            pending_by_hn = {}

            for row in pending_rows:
                hn = str(row["hn"] or "").strip()
                if not hn:
                    continue

                # HN เดิมซ้ำใน Queue -> รวมยาเข้ารายการแรกแล้วลบรายการซ้ำ
                if hn in pending_by_hn:
                    keep = pending_by_hn[hn]
                    try:
                        base = json.loads(keep["patient_json"] or "{}")
                    except Exception:
                        base = {}
                    try:
                        extra = json.loads(row["patient_json"] or "{}")
                    except Exception:
                        extra = {}

                    base.setdefault("medicines", [])
                    base["medicines"].extend(extra.get("medicines", []) or [])
                    if not base.get("dispense_date"):
                        base["dispense_date"] = row["dispense_date"] or extra.get("dispense_date", "")
                    if not base.get("appointment_date"):
                        base["appointment_date"] = row["appointment_date"] or extra.get("appointment_date", "")
                    base["required_days"] = calculate_days_between(
                        base.get("dispense_date", ""),
                        base.get("appointment_date", "")
                    )

                    conn.execute("""
                        UPDATE prescription_queue
                        SET patient_json = ?, dispense_date = ?, appointment_date = ?
                        WHERE id = ?
                    """, (
                        json.dumps(base, ensure_ascii=False, default=str),
                        str(base.get("dispense_date", "") or "").strip(),
                        str(base.get("appointment_date", "") or "").strip(),
                        keep["id"]
                    ))
                    conn.execute("DELETE FROM prescription_queue WHERE id = ?", (row["id"],))
                    continue

                pending_by_hn[hn] = row

            # ลบ/อัปเดตเฉพาะ pending ที่ตรงกับ Excel ปัจจุบัน
            for hn, row in list(pending_by_hn.items()):
                if hn not in source_map:
                    conn.execute(
                        "DELETE FROM prescription_queue WHERE id = ? AND status = 'pending'",
                        (row["id"],)
                    )
                    continue

                fresh_patient = source_map[hn]
                try:
                    current_patient = json.loads(row["patient_json"] or "{}")
                except Exception:
                    current_patient = {}

                if not current_patient.get("_manual_days_decision"):
                    conn.execute("""
                        UPDATE prescription_queue
                        SET dispense_date = ?, appointment_date = ?, patient_json = ?
                        WHERE id = ? AND status = 'pending'
                    """, (
                        str(fresh_patient.get("dispense_date", "") or "").strip(),
                        str(fresh_patient.get("appointment_date", "") or "").strip(),
                        json.dumps(fresh_patient, ensure_ascii=False, default=str),
                        row["id"]
                    ))

            # เพิ่ม HN ใหม่ที่ยังไม่มีใน Queue
            for hn, patient in source_map.items():
                if hn in pending_by_hn:
                    continue

                conn.execute("""
                    INSERT INTO prescription_queue (
                        hn, dispense_date, appointment_date, patient_json,
                        status, created_at
                    )
                    VALUES (?, ?, ?, ?, 'pending', ?)
                """, (
                    hn,
                    str(patient.get("dispense_date", "") or "").strip(),
                    str(patient.get("appointment_date", "") or "").strip(),
                    json.dumps(patient, ensure_ascii=False, default=str),
                    datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                ))

            conn.commit()
            return True
        finally:
            conn.close()

    except Exception as e:
        print("ERROR SYNC PRESCRIPTION QUEUE:", repr(e))
        return False


# ============================================================
# GET PENDING PATIENTS
# ============================================================

def get_pending_patients():

    # อ่าน Excel ล่าสุดทุกครั้งที่เรียก Queue
    sync_prescription_queue_from_excel()

    conn = get_db()


    try:

        rows = conn.execute(
            """
            SELECT
                id,
                hn,
                dispense_date,
                appointment_date,
                patient_json,
                created_at
            FROM prescription_queue
            WHERE status = 'pending'
            ORDER BY id ASC
            """
        ).fetchall()

    finally:

        conn.close()


    patients = []


    for row in rows:

        try:

            patient = json.loads(
                row["patient_json"]
            )

        except Exception:

            continue


        patient["queue_id"] = row["id"]

        patient["already_dispensed"] = False

        # ใช้วันของ Queue รายการนี้โดยตรง
        patient["hn"] = str(row["hn"] or patient.get("hn", "")).strip()
        patient["dispense_date"] = str(row["dispense_date"] or patient.get("dispense_date", "")).strip()
        patient["appointment_date"] = str(row["appointment_date"] or patient.get("appointment_date", "")).strip()
        patient["required_days"] = calculate_days_between(
            patient.get("dispense_date", ""),
            patient.get("appointment_date", "")
        )

        # ====================================================
        # CHECK INTERACTION AGAIN
        # ====================================================

        drug_names = [

            medicine.get(
                "name",
                ""
            )

            for medicine
            in patient.get(
                "medicines",
                []
            )

        ]


        # ====================================================
        # CALCULATE DAYS FIRST
        # ====================================================
        # คำนวณวันที่จ่าย -> วันนัดของ HN นี้ก่อนทุกครั้ง
        # เพื่อให้ตรวจสอบจำนวนยา/ช่วงวันของ HN นี้เสร็จก่อน
        patient = calculate_days_check(patient)

        # ====================================================
        # INTERACTION CHECK AFTER DAYS CHECK
        # ====================================================
        patient["interaction_results"] = check_patient_drug_interactions(
            patient
        )

        # สร้าง Consult PDF อัตโนมัติเมื่อจำนวนยาไม่สัมพันธ์กับวันนัด
        patient = ensure_days_consult_pdf(
            patient,
            queue_id=row["id"]
        )


        # ====================================================
        # CHECK CURRENT STOCK
        # ====================================================

        patient = apply_stock_to_patient(
            patient
        )


        patients.append(
            patient
        )


    # ========================================================
    # OLD → NEW
    # ========================================================

    patients.sort(

        key=lambda p:
        date_sort_key(
            p.get(
                "dispense_date",
                ""
            )
        )

    )


    return patients


# ============================================================
# GET ONE QUEUE PATIENT
# ============================================================

def get_queue_patient(
    queue_id
):

    conn = get_db()


    try:

        row = conn.execute(
            """
            SELECT *
            FROM prescription_queue
            WHERE id = ?
            AND status = 'pending'
            LIMIT 1
            """,
            (
                queue_id,
            )
        ).fetchone()

    finally:

        conn.close()


    if row is None:

        return None


    try:

        patient = json.loads(
            row["patient_json"]
        )

    except Exception:

        return None


    patient["queue_id"] = row["id"]

    # ใช้วันที่จาก Queue ของ HN นี้โดยตรง
    patient["hn"] = str(row["hn"] or patient.get("hn", "")).strip()
    patient["dispense_date"] = str(row["dispense_date"] or patient.get("dispense_date", "")).strip()
    patient["appointment_date"] = str(row["appointment_date"] or patient.get("appointment_date", "")).strip()
    patient["required_days"] = calculate_days_between(
        patient.get("dispense_date", ""),
        patient.get("appointment_date", "")
    )


    # ========================================================
    # CHECK INTERACTION
    # ========================================================

    drug_names = [

        medicine.get(
            "name",
            ""
        )

        for medicine
        in patient.get(
            "medicines",
            []
        )

    ]


    # ========================================================
    # CALCULATE DAYS FIRST
    # ========================================================
    # ใช้ dispense_date / appointment_date ของ HN นี้โดยตรง
    patient = calculate_days_check(patient)

    # ========================================================
    # CHECK INTERACTION AFTER DAYS CHECK
    # ========================================================
    patient["interaction_results"] = check_patient_drug_interactions(
        patient
    )

    # สร้าง Consult PDF อัตโนมัติถ้ายังไม่มี
    patient = ensure_days_consult_pdf(
        patient,
        queue_id=queue_id
    )


    # ========================================================
    # CHECK CURRENT STOCK
    # ========================================================

    patient = apply_stock_to_patient(
        patient
    )


    return patient


# ============================================================
# CHECK ALL STOCK SUFFICIENT
# ============================================================

def patient_stock_is_sufficient(
    patient
):

    stock_results = check_patient_stock(
        patient
    )


    if not stock_results:

        return False


    for item in stock_results:

        if not item.get(
            "found",
            False
        ):

            return False


        if item.get(
            "status"
        ) != "มีเพียงพอ":

            return False


    return True


# ============================================================
# COMPLETE QUEUE ACTION
# ============================================================

def complete_queue_action(
    patient,
    decision,
    decision_text,
    substitute_drug="",
    queue_id=None
):

    if queue_id is None:

        return False


    # ========================================================
    # PATIENT INFORMATION
    # ========================================================

    hn = str(
        patient.get(
            "hn",
            ""
        )
    ).strip()


    original_drugs = ", ".join(

        [

            str(
                medicine.get(
                    "name",
                    ""
                )
            )

            for medicine
            in patient.get(
                "medicines",
                []
            )

        ]

    )


    interaction_count = len(

        patient.get(
            "interaction_results",
            []
        )

    )


    # ========================================================
    # STOCK MUST BE DEDUCTED?
    #
    # dispense = จ่ายยาปกติ
    # original = Doctor ยืนยันใช้ยาเดิม
    # override_doctor = ผู้จ่ายยืนยันแทนแพทย์ตามใบสั่งเดิม
    #
    # substitute / cancel = ไม่หักในระบบปัจจุบัน
    # ========================================================

    # หัก Stock เฉพาะการจ่ายยา/ยืนยันใช้ยาเดิม
    # cancel และ substitute จะไม่หัก Stock
    should_deduct_stock = decision in {"dispense", "original", "override_doctor"}


    # ========================================================
    # CREATE STOCK BACKUP
    # ========================================================

    stock_backup_path = None


    if should_deduct_stock:

        if not os.path.exists(
            STOCK_FILE
        ):

            print(
                "ERROR: ไม่พบ stock.xlsx"
            )

            return False


        stock_backup_path = (
            STOCK_FILE
            +
            ".backup"
        )


        try:

            shutil.copyfile(

                STOCK_FILE,

                stock_backup_path

            )


        except Exception as e:

            print("")
            print("=" * 70)
            print("ERROR CREATE STOCK BACKUP")
            print(
                repr(e)
            )
            print("=" * 70)

            return False


    # ========================================================
    # DATABASE
    # ========================================================

    conn = get_db()


    try:

        # ====================================================
        # CHECK QUEUE
        # ====================================================

        queue_row = conn.execute(

            """
            SELECT id
            FROM prescription_queue
            WHERE id = ?
            AND status = 'pending'
            LIMIT 1
            """,

            (
                queue_id,
            )

        ).fetchone()


        if queue_row is None:

            conn.rollback()

            return False


        # ====================================================
        # CHECK HISTORY
        # ====================================================

        existing = conn.execute(

            """
            SELECT id
            FROM medication_history
            WHERE queue_id = ?
            LIMIT 1
            """,

            (
                queue_id,
            )

        ).fetchone()


        if existing is not None:

            conn.rollback()

            return False


        # ====================================================
        # DEDUCT STOCK
        # ====================================================

        if should_deduct_stock:

            stock_deducted = (
                deduct_stock_from_excel(
                    patient
                )
            )


            if not stock_deducted:

                conn.rollback()


                # --------------------------------------------
                # RESTORE STOCK
                # --------------------------------------------

                if (
                    stock_backup_path
                    and
                    os.path.exists(
                        stock_backup_path
                    )
                ):

                    try:

                        shutil.copyfile(

                            stock_backup_path,

                            STOCK_FILE

                        )

                    except Exception as restore_error:

                        print(
                            "ERROR RESTORE STOCK:",
                            repr(
                                restore_error
                            )
                        )


                return False


        # ====================================================
        # CURRENT TIME
        # ====================================================

        current_time = datetime.now().strftime(
            "%Y-%m-%d %H:%M:%S"
        )


        # ====================================================
        # INSERT HISTORY
        # ====================================================

        conn.execute(

            """
            INSERT INTO medication_history (
                queue_id,
                hn,
                patient_name,
                age,
                dispense_date,
                appointment_date,
                decision,
                decision_text,
                original_drugs,
                substitute_drug,
                interaction_count,
                dispensed_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,

            (

                queue_id,

                hn,

                patient.get(
                    "name",
                    ""
                ),

                patient.get(
                    "age",
                    ""
                ),

                patient.get(
                    "dispense_date",
                    ""
                ),

                patient.get(
                    "appointment_date",
                    ""
                ),

                decision,

                decision_text,

                original_drugs,

                substitute_drug,

                interaction_count,

                current_time

            )

        )


        # ====================================================
        # RESOLVE QUEUE
        # ====================================================

        cursor = conn.execute(

            """
            UPDATE prescription_queue
            SET
                status = 'resolved',
                resolved_at = ?
            WHERE id = ?
            AND status = 'pending'
            """,

            (

                current_time,

                queue_id

            )

        )


        # ====================================================
        # CHECK QUEUE UPDATE
        # ====================================================

        if cursor.rowcount != 1:

            conn.rollback()


            # ----------------------------------------------
            # RESTORE STOCK
            # ----------------------------------------------

            if (
                stock_backup_path
                and
                os.path.exists(
                    stock_backup_path
                )
            ):

                try:

                    shutil.copyfile(

                        stock_backup_path,

                        STOCK_FILE

                    )

                except Exception as restore_error:

                    print(
                        "ERROR RESTORE STOCK:",
                        repr(
                            restore_error
                        )
                    )


            return False


        # ====================================================
        # COMMIT
        # ====================================================

        conn.commit()


        # ====================================================
        # DELETE BACKUP
        # ====================================================

        if stock_backup_path:

            try:

                if os.path.exists(
                    stock_backup_path
                ):

                    os.remove(
                        stock_backup_path
                    )

            except Exception as e:

                print(
                    "WARNING DELETE BACKUP:",
                    repr(e)
                )


        # ====================================================
        # SUCCESS LOG
        # ====================================================

        print("")
        print("=" * 70)
        print("QUEUE COMPLETED SUCCESSFULLY")
        print(
            "Queue ID:",
            queue_id
        )
        print(
            "HN:",
            hn
        )
        print(
            "Decision:",
            decision
        )
        print(
            "Dispensed At:",
            current_time
        )
        print("=" * 70)


        return True


    # ========================================================
    # SQLITE INTEGRITY ERROR
    # ========================================================

    except sqlite3.IntegrityError as e:

        conn.rollback()


        # ====================================================
        # RESTORE STOCK
        # ====================================================

        if (
            stock_backup_path
            and
            os.path.exists(
                stock_backup_path
            )
        ):

            try:

                shutil.copyfile(

                    stock_backup_path,

                    STOCK_FILE

                )

            except Exception as restore_error:

                print(
                    "ERROR RESTORE STOCK:",
                    repr(
                        restore_error
                    )
                )


        print(
            "ERROR COMPLETE QUEUE:",
            repr(e)
        )


        return False


    # ========================================================
    # GENERAL ERROR
    # ========================================================

    except Exception as e:

        conn.rollback()


        # ====================================================
        # RESTORE STOCK
        # ====================================================

        if (
            stock_backup_path
            and
            os.path.exists(
                stock_backup_path
            )
        ):

            try:

                shutil.copyfile(

                    stock_backup_path,

                    STOCK_FILE

                )

            except Exception as restore_error:

                print(
                    "ERROR RESTORE STOCK:",
                    repr(
                        restore_error
                    )
                )


        print(
            "ERROR COMPLETE QUEUE:",
            repr(e)
        )


        return False


    finally:

        conn.close()


# ============================================================
# OLD RESOLVE QUEUE
# ============================================================

def resolve_queue(
    queue_id
):

    conn = get_db()


    try:

        cursor = conn.execute(
            """
            UPDATE prescription_queue
            SET
                status = 'resolved',
                resolved_at = ?
            WHERE id = ?
            AND status = 'pending'
            """,
            (
                datetime.now().strftime(
                    "%Y-%m-%d %H:%M:%S"
                ),

                queue_id
            )
        )


        conn.commit()


        return cursor.rowcount > 0


    except Exception as e:

        conn.rollback()

        print(
            "ERROR RESOLVE QUEUE:",
            repr(e)
        )

        return False


    finally:

        conn.close()


# ============================================================
# SAVE HISTORY
# ============================================================

def save_history(
    patient,
    decision,
    decision_text,
    substitute_drug="",
    queue_id=None
):

    if queue_id is None:

        return False


    hn = str(
        patient.get(
            "hn",
            ""
        )
    )


    original_drugs = ", ".join(

        [

            str(
                medicine.get(
                    "name",
                    ""
                )
            )

            for medicine
            in patient.get(
                "medicines",
                []
            )

        ]

    )


    interaction_count = len(

        patient.get(
            "interaction_results",
            []
        )

    )


    conn = get_db()


    try:

        existing = conn.execute(
            """
            SELECT id
            FROM medication_history
            WHERE queue_id = ?
            LIMIT 1
            """,
            (
                queue_id,
            )
        ).fetchone()


        if existing is not None:

            return False


        conn.execute(
            """
            INSERT INTO medication_history (
                queue_id,
                hn,
                patient_name,
                age,
                dispense_date,
                appointment_date,
                decision,
                decision_text,
                original_drugs,
                substitute_drug,
                interaction_count,
                dispensed_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                queue_id,

                hn,

                patient.get(
                    "name",
                    ""
                ),

                patient.get(
                    "age",
                    ""
                ),

                patient.get(
                    "dispense_date",
                    ""
                ),

                patient.get(
                    "appointment_date",
                    ""
                ),

                decision,

                decision_text,

                original_drugs,

                substitute_drug,

                interaction_count,

                datetime.now().strftime(
                    "%Y-%m-%d %H:%M:%S"
                )
            )
        )


        conn.commit()


        return True


    except Exception as e:

        conn.rollback()

        print(
            "ERROR SAVE HISTORY:",
            repr(e)
        )

        return False


    finally:

        conn.close()


# ============================================================
# GET HISTORY
# ============================================================

def get_history():

    conn = get_db()


    try:

        rows = conn.execute(
            """
            SELECT *
            FROM medication_history
            ORDER BY id DESC
            """
        ).fetchall()

    finally:

        conn.close()


    return rows


# ============================================================
# THAI FONT
# ============================================================

def register_thai_fonts():

    normal_font = None

    bold_font = None


    possible_normal_fonts = [

        r"C:\Windows\Fonts\THSarabunNew.ttf",

        r"C:\Windows\Fonts\THSarabunNew_Regular.ttf",

        r"C:\Windows\Fonts\THSarabunNew-Regular.ttf",

        r"C:\Windows\Fonts\tahoma.ttf",

        r"C:\Windows\Fonts\arial.ttf",

        r"C:\Windows\Fonts\LeelawUI.ttf"

    ]


    possible_bold_fonts = [

        r"C:\Windows\Fonts\THSarabunNew-Bold.ttf",

        r"C:\Windows\Fonts\THSarabunNew_Bold.ttf",

        r"C:\Windows\Fonts\tahomabd.ttf",

        r"C:\Windows\Fonts\arialbd.ttf",

        r"C:\Windows\Fonts\LeelaUIb.ttf"

    ]


    for font_path in possible_normal_fonts:

        if not os.path.exists(
            font_path
        ):

            continue


        try:

            pdfmetrics.registerFont(
                TTFont(
                    "ThaiNormal",
                    font_path
                )
            )


            normal_font = "ThaiNormal"

            break


        except Exception:

            pass


    for font_path in possible_bold_fonts:

        if not os.path.exists(
            font_path
        ):

            continue


        try:

            pdfmetrics.registerFont(
                TTFont(
                    "ThaiBold",
                    font_path
                )
            )


            bold_font = "ThaiBold"

            break


        except Exception:

            pass


    if normal_font is None:

        normal_font = "Helvetica"


    if bold_font is None:

        bold_font = normal_font


    return normal_font, bold_font


# ============================================================
# DRAW WRAPPED TEXT
# ============================================================

def draw_wrapped_text(
    pdf,
    text,
    x,
    y,
    max_chars=70,
    font="ThaiNormal",
    size=14,
    line_height=18
):

    if text is None:

        text = ""


    text = str(
        text
    )


    lines = []


    while len(text) > max_chars:

        cut = text.rfind(
            " ",
            0,
            max_chars
        )


        if cut <= 0:

            cut = max_chars


        line = text[
            :cut
        ].strip()


        if line:

            lines.append(
                line
            )


        text = text[
            cut:
        ].strip()


    if text:

        lines.append(
            text
        )


    pdf.setFont(
        font,
        size
    )


    for line in lines:

        pdf.drawString(
            x,
            y,
            line
        )


        y -= line_height


    return y


# ============================================================
# SAVE UPDATED QUEUE PATIENT
# ============================================================

def update_queue_patient(queue_id, patient):
    conn = get_db()

    try:
        conn.execute(
            """
            UPDATE prescription_queue
            SET patient_json = ?
            WHERE id = ? AND status = 'pending'
            """,
            (
                json.dumps(
                    patient,
                    ensure_ascii=False,
                    default=str
                ),
                queue_id
            )
        )
        conn.commit()
        return True

    except Exception as e:
        conn.rollback()
        print(
            "ERROR UPDATE QUEUE PATIENT:",
            repr(e)
        )
        return False

    finally:
        conn.close()


# ============================================================
# CREATE DAYS SUPPLY CONSULT PDF
# ============================================================

def create_days_supply_consult_pdf(patient):
    """
    สร้าง Days Supply Consult PDF จากข้อมูล patient ของ HN เดียวกันโดยตรง
    แสดงยาทุกตัวใน patient["medicines"] พร้อมผลคำนวณ Days Supply
    ไม่ใช้ days_check_results เป็นแหล่งข้อมูลหลัก เพื่อป้องกันข้อมูลเก่าหรือ key ไม่ตรงกัน
    """
    try:
        hn = str(patient.get("hn", "") or "").strip()
        patient_name = str(patient.get("name", "") or "").strip() or "ไม่ได้ระบุ"
        age = str(patient.get("age", "") or "").strip() or "ไม่ได้ระบุ"
        dispense_date = str(patient.get("dispense_date", "") or "").strip() or "ไม่ได้ระบุ"
        appointment_date = str(patient.get("appointment_date", "") or "").strip() or "ไม่ได้ระบุ"

        # คำนวณใหม่จากวันที่ของ HN นี้โดยตรง
        required_days = calculate_days_between(
            patient.get("dispense_date", ""),
            patient.get("appointment_date", "")
        )
        patient["required_days"] = required_days

        # ใช้ยาจริงจาก patient เป็นแหล่งข้อมูลหลัก
        medicines = patient.get("medicines", []) or []

        # ถ้ามี days_check_results เก่าค้างอยู่ ให้ map ผลตามชื่อยาไว้เป็นข้อมูลเสริมเท่านั้น
        result_map = {}
        for result in patient.get("days_check_results", []) or []:
            key = normalize_drug_name(result.get("name", ""))
            if key:
                result_map[key] = result

        filename = "Days_Supply_Consult_HN_{}.pdf".format(
            hn.replace(" ", "_") or "UNKNOWN"
        )
        pdf_path = os.path.join(CONSULT_FOLDER, filename)

        normal_font, bold_font = register_thai_fonts()
        pdf = canvas.Canvas(pdf_path, pagesize=A4)
        width, height = A4
        y = height - 50

        def page_space(current_y, required=70):
            if current_y < required:
                pdf.showPage()
                return height - 50
            return current_y

        def draw_field(label, value):
            nonlocal y
            y = page_space(y, 55)
            y = draw_wrapped_text(
                pdf,
                "{}: {}".format(label, value),
                60, y, 75, normal_font, 14, 19
            )

        pdf.setFont(bold_font, 22)
        pdf.drawCentredString(width / 2, y, "MEDICATION DAYS SUPPLY CONSULT")
        y -= 30
        pdf.setFont(bold_font, 18)
        pdf.drawCentredString(
            width / 2, y,
            "บันทึกปรึกษาความสัมพันธ์ระหว่างจำนวนยาและวันนัด"
        )
        y -= 45

        draw_field("HN", hn or "ไม่ได้ระบุ")
        draw_field("ชื่อผู้ป่วย", patient_name)
        draw_field("อายุ", age)
        draw_field("วันที่จ่ายยา", dispense_date)
        draw_field("วันนัด", appointment_date)
        draw_field("จำนวนวันที่ต้องใช้", format_quantity(required_days) + " วัน")

        y -= 15
        y = page_space(y, 100)
        pdf.setFont(bold_font, 17)
        pdf.drawString(50, y, "รายการยาที่ตรวจสอบ")
        y -= 28

        if not medicines:
            y = page_space(y, 60)
            y = draw_wrapped_text(
                pdf,
                "ไม่พบรายการยาในข้อมูลผู้ป่วย",
                70, y, 70, normal_font, 14, 18
            )
        else:
            for index, medicine in enumerate(medicines):
                y = page_space(y, 190)

                name = str(medicine.get("name", "") or "").strip() or "ไม่ได้ระบุ"
                strength = str(medicine.get("strength", "") or "").strip() or "ไม่ได้ระบุ"

                try:
                    times = float(medicine.get("times_per_day", 0) or 0)
                except Exception:
                    times = 0

                try:
                    quantity = float(medicine.get("quantity", 0) or 0)
                except Exception:
                    quantity = 0

                # คำนวณจากข้อมูลยาจริงทุกครั้ง
                expected = required_days * times if required_days > 0 and times > 0 else 0
                missing = max(expected - quantity, 0)
                excess = max(quantity - expected, 0)

                # ใช้ค่าที่คำนวณสดเป็นหลัก และเก็บค่าลง medicine ด้วย
                medicine["required_days"] = required_days
                medicine["expected_quantity"] = int(expected) if float(expected).is_integer() else expected
                medicine["missing_quantity"] = int(missing) if float(missing).is_integer() else missing
                medicine["excess_quantity"] = int(excess) if float(excess).is_integer() else excess

                if required_days <= 0 or times <= 0:
                    status = "คำนวณไม่ได้ — กรุณาตรวจสอบวันที่จ่าย/วันนัด/ครั้งต่อวัน"
                elif missing > 0:
                    status = "ขาด {} เม็ด".format(format_quantity(missing))
                elif excess > 0:
                    status = "เกิน {} เม็ด".format(format_quantity(excess))
                else:
                    status = "พอดี"

                pdf.setFont(bold_font, 16)
                pdf.drawString(60, y, "รายการยา {}: {}".format(index + 1, name))
                y -= 24

                fields = [
                    "ขนาดยา: {}".format(str(strength)),
                    "ครั้งต่อวัน: {} ครั้ง".format(
                        format_quantity(times) if times > 0 else "ไม่ได้ระบุ"
                    ),
                    "จำนวนวันที่ต้องใช้: {} วัน".format(format_quantity(required_days)),
                    "จำนวนยาที่สั่ง: {} เม็ด".format(format_quantity(quantity)),
                    "จำนวนยาที่ควรเป็น: {} เม็ด".format(format_quantity(expected)),
                    "ผลการตรวจ: {}".format(status)
                ]

                for field in fields:
                    y = page_space(y, 55)
                    y = draw_wrapped_text(
                        pdf, field, 70, y, 70, normal_font, 14, 18
                    )

                # ถ้าต้องปรึกษา ให้แสดงรายการที่มีปัญหาอย่างชัดเจน
                if missing > 0 or excess > 0:
                    y = page_space(y, 55)
                    problem_text = (
                        "*** ขาด {} เม็ด ***".format(format_quantity(missing))
                        if missing > 0 else
                        "*** เกิน {} เม็ด ***".format(format_quantity(excess))
                    )
                    y = draw_wrapped_text(
                        pdf, problem_text, 70, y, 70, bold_font, 14, 18
                    )

                y -= 12

        # Interaction: ใช้ผลของ HN นี้เท่านั้น
        y = page_space(y, 150)
        y -= 5
        pdf.setFont(bold_font, 17)
        pdf.drawString(50, y, "ยาที่มีปฏิกิริยาระหว่างกัน")
        y -= 28

        interaction_results = patient.get("interaction_results", []) or []
        if interaction_results:
            for index, interaction in enumerate(interaction_results, 1):
                y = page_space(y, 100)
                drug1 = str(interaction.get("Drug_1", interaction.get("drug1", "")) or "").strip()
                drug2 = str(interaction.get("Drug_2", interaction.get("drug2", "")) or "").strip()
                risk = str(interaction.get("Risk", interaction.get("risk", "")) or "ไม่ได้ระบุ")
                severity = str(interaction.get("Severity", interaction.get("severity", "")) or "ไม่ได้ระบุ")
                significance = str(
                    interaction.get(
                        "Clinical_Significance",
                        interaction.get("clinical_significance", "")
                    ) or "ไม่ได้ระบุ"
                )
                reference = str(
                    interaction.get("Reference", interaction.get("reference", "")) or "ไม่ได้ระบุ"
                )

                for field in [
                    "{}. {} ↔ {}".format(index, drug1 or "ไม่ได้ระบุ", drug2 or "ไม่ได้ระบุ"),
                    "Risk: {}".format(risk),
                    "Severity: {}".format(severity),
                    "Clinical Significance: {}".format(significance),
                    "Reference: {}".format(reference)
                ]:
                    y = page_space(y, 55)
                    y = draw_wrapped_text(
                        pdf, field, 65, y, 70, normal_font, 13, 17
                    )
                y -= 8
        else:
            y = draw_wrapped_text(
                pdf,
                "ไม่พบยาที่มีปฏิกิริยาระหว่างกัน",
                65, y, 70, normal_font, 14, 18
            )

        y = page_space(y, 190)
        y -= 10
        pdf.setFont(bold_font, 17)
        pdf.drawString(50, y, "ความเห็นแพทย์")
        y -= 30
        pdf.setFont(normal_font, 14)

        for text_line in [
            "☐ เพิ่มจำนวนยาให้ครบตามจำนวนที่ขาด",
            "☐ ลดจำนวนยาในส่วนที่เกิน",
            "☐ เห็นควรจ่ายตามจำนวนเดิม",
            "☐ เห็นควรจ่ายยาตามเดิม",
            "☐ ปรับเปลี่ยน/หยุดยาที่มีปฏิกิริยาระหว่างกัน",
            "☐ ติดตามอาการหรือผลตรวจทางห้องปฏิบัติการเพิ่มเติม",
            "☐ อื่น ๆ: ________________________________________________"
        ]:
            y = page_space(y, 60)
            y = draw_wrapped_text(
                pdf, text_line, 65, y, 70, normal_font, 14, 22
            )

        y -= 15
        for text_line in [
            "แพทย์ผู้พิจารณา: ______________________________",
            "วันที่: _________________________________________",
            "หมายเหตุ: _____________________________________"
        ]:
            y = page_space(y, 50)
            pdf.setFont(normal_font, 14)
            pdf.drawString(65, y, text_line)
            y -= 25

        pdf.setFont(normal_font, 9)
        pdf.drawCentredString(
            width / 2, 25,
            "ระบบ Medication Management / Clinical Decision Support"
        )
        pdf.save()

        if not os.path.exists(pdf_path):
            return None

        return filename

    except Exception as e:
        print("ERROR CREATE DAYS SUPPLY CONSULT PDF:", repr(e))
        return None


# ============================================================
# ENSURE DAYS SUPPLY CONSULT PDF
# ============================================================

def ensure_days_consult_pdf(patient, queue_id=None):
    """
    ตรวจว่าผู้ป่วยมีปัญหา Days Supply และทำให้ Consult PDF เป็นฉบับล่าสุดเสมอ
    ถ้าเป็น PDF เก่าจากโครงสร้างข้อมูลเดิม จะสร้างใหม่อัตโนมัติ
    """
    if not patient.get("days_check_required", False):
        return patient

    # สร้างลายเซ็นจากข้อมูลที่ PDF ใช้จริง เพื่อไม่สร้างไฟล์ใหม่ทุก refresh
    signature_data = {
        "hn": patient.get("hn", ""),
        "name": patient.get("name", ""),
        "age": patient.get("age", ""),
        "dispense_date": patient.get("dispense_date", ""),
        "appointment_date": patient.get("appointment_date", ""),
        "required_days": patient.get("required_days", 0),
        "medicines": patient.get("medicines", []) or [],
        "interaction_results": patient.get("interaction_results", []) or []
    }
    signature = hashlib.sha256(
        json.dumps(signature_data, ensure_ascii=False, sort_keys=True, default=str).encode("utf-8")
    ).hexdigest()

    existing_filename = str(patient.get("days_consult_pdf", "") or "").strip()
    existing_signature = str(patient.get("days_consult_signature", "") or "").strip()

    if (
        existing_filename
        and existing_signature == signature
        and os.path.exists(os.path.join(CONSULT_FOLDER, existing_filename))
    ):
        return patient

    # ลบ PDF เก่าของ HN เดียวกันทั้งหมด เพื่อไม่ให้ดาวน์โหลดไฟล์เก่าที่ข้อมูลไม่ครบ
    hn = str(patient.get("hn", "") or "").strip().replace(" ", "_")
    if hn:
        prefix = "Days_Supply_Consult_HN_{}_".format(hn)
        try:
            for old_name in os.listdir(CONSULT_FOLDER):
                if old_name.startswith(prefix) and old_name.lower().endswith(".pdf"):
                    try:
                        os.remove(os.path.join(CONSULT_FOLDER, old_name))
                    except OSError:
                        pass
        except Exception:
            pass

    filename = create_days_supply_consult_pdf(patient)

    if filename:
        patient["days_consult_pdf"] = filename
        patient["days_consult_signature"] = signature
        patient["days_consult_created_at"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        if queue_id is not None:
            update_queue_patient(queue_id, patient)

    return patient


# ============================================================
# SAVE DAYS CONSULT PDF NAME
# ============================================================

def save_days_consult_pdf_name(queue_id, patient, filename):
    patient["days_consult_pdf"] = filename
    patient["days_consult_created_at"] = (
        datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    )
    return update_queue_patient(
        queue_id,
        patient
    )


# ============================================================
# CREATE CONSULT PDF
# ============================================================

@app.route(
    "/create-consult-pdf",
    methods=["POST"]
)
def create_consult_pdf():

    try:

        hn = request.form.get(
            "hn",
            ""
        )


        patient_name = request.form.get(
            "patient_name",
            ""
        )


        age = request.form.get(
            "age",
            ""
        )


        dispense_date = request.form.get(
            "dispense_date",
            ""
        )


        appointment_date = request.form.get(
            "appointment_date",
            ""
        )


        ward = request.form.get(
            "ward",
            "ไม่ได้ระบุในข้อมูลใบสั่งยา"
        )


        diagnosis = request.form.get(
            "diagnosis",
            "ไม่ได้ระบุในข้อมูลใบสั่งยา"
        )


        weight = request.form.get(
            "weight",
            "ไม่ได้ระบุในข้อมูลใบสั่งยา"
        )


        height = request.form.get(
            "height",
            "ไม่ได้ระบุในข้อมูลใบสั่งยา"
        )


        lab = request.form.get(
            "lab",
            "ไม่ได้ระบุในข้อมูลใบสั่งยา"
        )


        renal_function = request.form.get(
            "renal_function",
            "ไม่ได้ระบุในข้อมูลใบสั่งยา"
        )


        hepatic_function = request.form.get(
            "hepatic_function",
            "ไม่ได้ระบุในข้อมูลใบสั่งยา"
        )


        try:

            interaction_count = int(
                request.form.get(
                    "interaction_count",
                    0
                )
            )

        except Exception:

            interaction_count = 0


        interactions = []


        for i in range(
            interaction_count
        ):

            interactions.append({

                "Drug_1":
                    request.form.get(
                        "drug1_{}".format(i),
                        ""
                    ),

                "Drug_2":
                    request.form.get(
                        "drug2_{}".format(i),
                        ""
                    ),

                "Risk":
                    request.form.get(
                        "risk_{}".format(i),
                        ""
                    ),

                "Severity":
                    request.form.get(
                        "severity_{}".format(i),
                        ""
                    ),

                "Summary":
                    request.form.get(
                        "summary_{}".format(i),
                        ""
                    ),

                "Management":
                    request.form.get(
                        "management_{}".format(i),
                        ""
                    ),

                "Reference":
                    request.form.get(
                        "reference_{}".format(i),
                        ""
                    )

            })


        if not interactions:

            return """
            <html lang="th">
            <head>
                <meta charset="UTF-8">
                <title>ไม่พบ Interaction</title>
            </head>
            <body>
                <h2>ไม่พบข้อมูล Drug Interaction</h2>
                <p>
                    ไม่สามารถสร้าง Consult PDF
                    เนื่องจากไม่มี Interaction
                </p>
                <a href="/prescription-result">
                    ← กลับไปตรวจใบสั่งยา
                </a>
            </body>
            </html>
            """


        # ====================================================
        # MEDICATIONS
        # ====================================================

        medications = []


        try:

            medication_count = int(
                request.form.get(
                    "medication_count",
                    0
                )
            )

        except Exception:

            medication_count = 0


        for i in range(
            medication_count
        ):

            medications.append({

                "name":
                    request.form.get(
                        "med_name_{}".format(i),
                        "ไม่ได้ระบุ"
                    ),

                "strength":
                    request.form.get(
                        "med_strength_{}".format(i),
                        "ไม่ได้ระบุ"
                    ),

                "frequency":
                    request.form.get(
                        "med_frequency_{}".format(i),
                        "ไม่ได้ระบุ"
                    ),

                "route":
                    request.form.get(
                        "med_route_{}".format(i),
                        "ไม่ได้ระบุ"
                    ),

                "start_date":
                    request.form.get(
                        "med_start_date_{}".format(i),
                        "ไม่ได้ระบุ"
                    ),

                "quantity":
                    request.form.get(
                        "med_quantity_{}".format(i),
                        "ไม่ได้ระบุ"
                    )

            })


        # ====================================================
        # FALLBACK
        # ====================================================

        if not medications:

            added_names = set()


            for item in interactions:

                for key in [
                    "Drug_1",
                    "Drug_2"
                ]:

                    drug = item.get(
                        key,
                        ""
                    )


                    if (
                        drug
                        and
                        drug not in added_names
                    ):

                        medications.append({

                            "name": drug,

                            "strength":
                                "ไม่ได้ระบุ",

                            "frequency":
                                "ไม่ได้ระบุ",

                            "route":
                                "ไม่ได้ระบุ",

                            "start_date":
                                "ไม่ได้ระบุ",

                            "quantity":
                                "ไม่ได้ระบุ"

                        })


                        added_names.add(
                            drug
                        )


        # ====================================================
        # FILE
        # ====================================================

        timestamp = datetime.now().strftime(
            "%Y%m%d_%H%M%S"
        )


        safe_hn = str(
            hn
        ).replace(
            " ",
            "_"
        )


        filename = (
            "Drug_Interaction_Consult_HN_{}_{}.pdf"
            .format(
                safe_hn,
                timestamp
            )
        )


        pdf_path = os.path.join(
            CONSULT_FOLDER,
            filename
        )


        normal_font, bold_font = (
            register_thai_fonts()
        )


        pdf = canvas.Canvas(
            pdf_path,
            pagesize=A4
        )


        width, height = A4


        def new_page():

            pdf.showPage()

            return height - 50


        y = height - 50


        # ====================================================
        # TITLE
        # ====================================================

        pdf.setFont(
            bold_font,
            22
        )


        pdf.drawCentredString(
            width / 2,
            y,
            "DRUG INTERACTION CONSULT"
        )


        y -= 30


        pdf.setFont(
            bold_font,
            18
        )


        pdf.drawCentredString(
            width / 2,
            y,
            "บันทึกปรึกษาปัญหายาระหว่างยา"
        )


        y -= 45


        # ====================================================
        # 1
        # ====================================================

        pdf.setFont(
            bold_font,
            17
        )


        pdf.drawString(
            50,
            y,
            "1. หัวข้อการปรึกษา"
        )


        y -= 28


        subject_text = (
            "ขอปรึกษาเรื่อง Drug Interaction "
            "ในผู้ป่วย HN {} ระหว่างยา {} และ {}"
        ).format(

            hn,

            interactions[0].get(
                "Drug_1",
                ""
            ),

            interactions[0].get(
                "Drug_2",
                ""
            )

        )


        y = draw_wrapped_text(
            pdf,
            subject_text,
            60,
            y,
            75,
            normal_font,
            15,
            20
        )


        y -= 10


        for text in [

            "HN: {}".format(
                hn
            ),

            "ชื่อผู้ป่วย: {}".format(
                patient_name
            ),

            "อายุ: {}".format(
                age
            ),

            "วอร์ด/แผนก: {}".format(
                ward
            )

        ]:

            y = draw_wrapped_text(
                pdf,
                text,
                60,
                y,
                75,
                normal_font,
                15,
                20
            )


        # ====================================================
        # 2
        # ====================================================

        if y < 180:

            y = new_page()


        y -= 15


        pdf.setFont(
            bold_font,
            17
        )


        pdf.drawString(
            50,
            y,
            "2. ข้อมูลทางคลินิกที่สำคัญ"
        )


        y -= 28


        clinical_items = [

            "Diagnosis หลัก: {}".format(
                diagnosis
            ),

            "น้ำหนัก: {}".format(
                weight
            ),

            "ส่วนสูง: {}".format(
                height
            ),

            "Lab ที่เกี่ยวข้อง: {}".format(
                lab
            ),

            "การทำงานของไต: {}".format(
                renal_function
            ),

            "การทำงานของตับ: {}".format(
                hepatic_function
            ),

            "วันที่จ่ายยา: {}".format(
                dispense_date
            ),

            "วันนัด: {}".format(
                appointment_date
            )

        ]


        for text in clinical_items:

            if y < 80:

                y = new_page()


            y = draw_wrapped_text(
                pdf,
                text,
                60,
                y,
                75,
                normal_font,
                14,
                19
            )


        # ====================================================
        # 3
        # ====================================================

        if y < 180:

            y = new_page()


        y -= 10


        pdf.setFont(
            bold_font,
            17
        )


        pdf.drawString(
            50,
            y,
            "3. รายการยาปัจจุบันของผู้ป่วย"
        )


        y -= 28


        for index, med in enumerate(
            medications
        ):

            if y < 150:

                y = new_page()


                pdf.setFont(
                    bold_font,
                    17
                )


                pdf.drawString(
                    50,
                    y,
                    "3. รายการยาปัจจุบันของผู้ป่วย (ต่อ)"
                )


                y -= 28


            y = draw_wrapped_text(
                pdf,
                "{}. {}".format(
                    index + 1,
                    med.get(
                        "name",
                        "ไม่ได้ระบุ"
                    )
                ),
                60,
                y,
                70,
                bold_font,
                15,
                20
            )


            for detail in [

                "ขนาดยา: {}".format(
                    med.get(
                        "strength",
                        "ไม่ได้ระบุ"
                    )
                ),

                "วิธีใช้/ความถี่: {}".format(
                    med.get(
                        "frequency",
                        "ไม่ได้ระบุ"
                    )
                ),

                "Route: {}".format(
                    med.get(
                        "route",
                        "ไม่ได้ระบุ"
                    )
                ),

                "วันที่เริ่มยา: {}".format(
                    med.get(
                        "start_date",
                        "ไม่ได้ระบุ"
                    )
                ),

                "จำนวนยา: {}".format(
                    med.get(
                        "quantity",
                        "ไม่ได้ระบุ"
                    )
                )

            ]:

                y = draw_wrapped_text(
                    pdf,
                    detail,
                    75,
                    y,
                    70,
                    normal_font,
                    14,
                    18
                )


            y -= 8


        # ====================================================
        # 4
        # ====================================================

        if y < 180:

            y = new_page()


        pdf.setFont(
            bold_font,
            17
        )


        pdf.drawString(
            50,
            y,
            "4. ปัญหาและปฏิกิริยาระหว่างยาที่พบ"
        )


        y -= 28


        for index, item in enumerate(
            interactions
        ):

            if y < 180:

                y = new_page()


                pdf.setFont(
                    bold_font,
                    17
                )


                pdf.drawString(
                    50,
                    y,
                    "4. ปัญหาและปฏิกิริยาระหว่างยา (ต่อ)"
                )


                y -= 28


            fields = [

                "Interaction {}".format(
                    index + 1
                ),

                "Drug 1: {}".format(
                    item.get(
                        "Drug_1",
                        ""
                    )
                ),

                "Drug 2: {}".format(
                    item.get(
                        "Drug_2",
                        ""
                    )
                ),

                "Risk: {}".format(
                    item.get(
                        "Risk",
                        "ไม่ได้ระบุ"
                    )
                ),

                "Severity: {}".format(
                    item.get(
                        "Severity",
                        "ไม่ได้ระบุ"
                    )
                ),

                "Clinical Significance: {}".format(
                    item.get(
                        "Summary",
                        "ไม่ได้ระบุ"
                    )
                ),

                "Reference: {}".format(
                    item.get(
                        "Reference",
                        "ไม่ได้ระบุ"
                    )
                )

            ]


            for field_index, text in enumerate(
                fields
            ):

                y = draw_wrapped_text(
                    pdf,
                    text,
                    60 if field_index == 0 else 70,
                    y,
                    70,
                    bold_font if field_index == 0 else normal_font,
                    14,
                    18
                )


            y -= 8


        # ====================================================
        # 5
        # ====================================================

        if y < 180:

            y = new_page()


        pdf.setFont(
            bold_font,
            17
        )


        pdf.drawString(
            50,
            y,
            "5. ข้อเสนอแนะทางเภสัชบำบัด"
        )


        y -= 28


        for index, item in enumerate(
            interactions
        ):

            management = item.get(
                "Management",
                "ไม่ได้ระบุ"
            )


            y = draw_wrapped_text(
                pdf,
                "Interaction {}: {}".format(
                    index + 1,
                    management
                ),
                60,
                y,
                70,
                normal_font,
                14,
                18
            )


            if y < 100:

                y = new_page()


        y -= 10


        y = draw_wrapped_text(
            pdf,
            "การติดตาม: ควรพิจารณาติดตามอาการไม่พึงประสงค์ "
            "และผลตรวจทางห้องปฏิบัติการที่เกี่ยวข้อง "
            "ตามความเหมาะสมของผู้ป่วย",
            60,
            y,
            70,
            normal_font,
            14,
            18
        )


        # ====================================================
        # 6
        # ====================================================

        if y < 180:

            y = new_page()


        y -= 15


        pdf.setFont(
            bold_font,
            17
        )


        pdf.drawString(
            50,
            y,
            "6. สรุปและลงชื่อผู้ปรึกษา"
        )


        y -= 28


        conclusion = (
            "พบข้อมูล Drug Interaction จำนวน {} รายการ "
            "จากฐานความรู้ของระบบ "
            "จึงขอปรึกษาแพทย์เพื่อพิจารณา "
            "ความเหมาะสมของการใช้ยาร่วมกัน"
        ).format(
            len(interactions)
        )


        y = draw_wrapped_text(
            pdf,
            conclusion,
            60,
            y,
            70,
            normal_font,
            14,
            18
        )


        y -= 25


        for text in [

            "ผู้ปรึกษา: ______________________________",

            "ตำแหน่ง: ภก. / นศ.ภ. ____________________",

            "วันที่: __________________________________",

            "ช่องทางติดต่อกลับ: _______________________"

        ]:

            if y < 60:

                y = new_page()


            pdf.setFont(
                normal_font,
                14
            )


            pdf.drawString(
                60,
                y,
                text
            )


            y -= 25


        # ====================================================
        # FOOTER
        # ====================================================

        pdf.setFont(
            normal_font,
            9
        )


        pdf.drawCentredString(
            width / 2,
            25,
            "ระบบ Medication Management / Clinical Decision Support"
        )


        pdf.save()


        if not os.path.exists(
            pdf_path
        ):

            raise Exception(
                "ไม่พบไฟล์ PDF หลังจากสร้าง"
            )


        return render_template(
            "consult_success.html",
            filename=filename
        )


    except Exception as e:

        print("=" * 70)
        print("ERROR CREATE CONSULT PDF")
        print(repr(e))
        print("=" * 70)


        return render_template(
            "consult_success.html",
            filename=None,
            error=str(e)
        )


# ============================================================
# DOWNLOAD CONSULT PDF
# ============================================================

@app.route(
    "/download-consult/<path:filename>"
)
def download_consult(filename):

    return send_from_directory(
        CONSULT_FOLDER,
        filename,
        as_attachment=True
    )


# ============================================================
# PRESCRIPTION UPLOAD
# ============================================================

@app.route(
    "/prescription",
    methods=["GET", "POST"]
)
def prescription():

    if request.method == "GET":

        return render_template(
            "prescription.html"
        )


    uploaded_file = request.files.get(
        "file"
    )


    if uploaded_file is None:

        return render_template(
            "prescription.html",
            error="กรุณาเลือกไฟล์ Excel"
        )


    if uploaded_file.filename == "":

        return render_template(
            "prescription.html",
            error="กรุณาเลือกไฟล์ Excel"
        )


    extension = os.path.splitext(
        uploaded_file.filename
    )[1].lower()


    if extension not in [
        ".xlsx",
        ".xlsm"
    ]:

        return render_template(
            "prescription.html",
            error="กรุณาเลือกไฟล์ Excel (.xlsx หรือ .xlsm)"
        )


    temp_path = None


    try:

        with tempfile.NamedTemporaryFile(
            delete=False,
            suffix=extension
        ) as temp_file:

            uploaded_file.save(
                temp_file.name
            )

            temp_path = temp_file.name


        # เก็บ Excel ต้นฉบับไว้บน Server เพื่อให้ Queue Sync ได้แบบ Real-time
        shutil.copyfile(
            temp_path,
            PRESCRIPTION_SOURCE_FILE
        )

        patient_results = read_prescription_excel(
            temp_path
        )


        if not patient_results:

            return render_template(
                "prescription.html",
                error="ไม่พบข้อมูลผู้ป่วยในไฟล์ Excel"
            )


        added_count = 0

        duplicate_count = 0


        for patient in patient_results:

            added = add_patient_to_queue(
                patient
            )


            if added:

                added_count += 1

            else:

                duplicate_count += 1


        print(
            "เพิ่มผู้ป่วยเข้า Queue:",
            added_count,
            "ราย"
        )


        print(
            "ข้อมูลซ้ำ/ไม่เพิ่ม:",
            duplicate_count,
            "ราย"
        )


        return redirect(
            url_for(
                "prescription_result"
            )
        )


    except Exception as e:

        print(
            "ERROR READING PRESCRIPTION EXCEL:",
            repr(e)
        )


        return render_template(
            "prescription.html",
            error=(
                "เกิดข้อผิดพลาดในการอ่าน Excel: "
                + str(e)
            )
        )


    finally:

        if temp_path:

            try:

                if os.path.exists(
                    temp_path
                ):

                    os.remove(
                        temp_path
                    )

            except Exception:

                pass


# ============================================================
# ADD PATIENT TO QUEUE
# ============================================================

def add_patient_to_queue(
    patient
):

    hn = str(
        patient.get(
            "hn",
            ""
        )
    ).strip()


    dispense_date = str(
        patient.get(
            "dispense_date",
            ""
        )
    ).strip()


    if not hn:

        return False


    conn = get_db()


    try:

        existing = conn.execute(
            """
            SELECT id
            FROM prescription_queue
            WHERE hn = ?
            AND status = 'pending'
            LIMIT 1
            """,
            (hn,)
        ).fetchone()


        if existing is not None:

            return False


        # ----------------------------------------------------
        # สร้าง Days Supply Consult ตั้งแต่ตอนเพิ่มเข้า Queue
        # เพื่อให้หน้าแสดงผลมีข้อมูลแจ้งเตือนและปุ่มดาวน์โหลดทันที
        # ----------------------------------------------------
        if patient.get("days_check_required", False):
            pdf_filename = create_days_supply_consult_pdf(patient)
            if pdf_filename:
                patient["days_consult_pdf"] = pdf_filename
                patient["days_consult_created_at"] = datetime.now().strftime(
                    "%Y-%m-%d %H:%M:%S"
                )

        conn.execute(
            """
            INSERT INTO prescription_queue (
                hn,
                dispense_date,
                appointment_date,
                patient_json,
                status,
                created_at
            )
            VALUES (?, ?, ?, ?, 'pending', ?)
            """,
            (
                hn,
                dispense_date,
                str(patient.get("appointment_date", "") or "").strip(),
                json.dumps(
                    patient,
                    ensure_ascii=False,
                    default=str
                ),
                datetime.now().strftime(
                    "%Y-%m-%d %H:%M:%S"
                )
            )
        )


        conn.commit()


        return True


    except sqlite3.IntegrityError:

        conn.rollback()

        return False


    except Exception as e:

        conn.rollback()

        print(
            "ERROR ADD QUEUE:",
            repr(e)
        )

        return False


    finally:

        conn.close()


# ============================================================
# PRESCRIPTION RESULT / QUEUE
# ============================================================

@app.route(
    "/prescription-result",
    methods=["GET"]
)
def prescription_result():

    patients = get_pending_patients()


    return render_template(
        "prescription_result.html",
        patients=patients
    )


# ============================================================
# REAL-TIME PRESCRIPTION REFRESH
# ============================================================

@app.route("/prescription-live")
def prescription_live():
    """
    Endpoint สำหรับหน้าเว็บตรวจว่า Excel เปลี่ยนหรือไม่
    """
    sync_prescription_queue_from_excel()

    if os.path.exists(PRESCRIPTION_SOURCE_FILE):
        try:
            stat = os.stat(PRESCRIPTION_SOURCE_FILE)
            signature = "{}-{}".format(
                stat.st_mtime_ns,
                stat.st_size
            )
        except Exception:
            signature = "unknown"
    else:
        signature = "missing"

    conn = get_db()
    try:
        pending_count = conn.execute(
            """
            SELECT COUNT(*)
            FROM prescription_queue
            WHERE status = 'pending'
            """
        ).fetchone()[0]
    finally:
        conn.close()

    return jsonify({
        "signature": signature,
        "pending_count": pending_count
    })


# ============================================================
# DISPENSE
# ============================================================

@app.route(
    "/dispense/<int:queue_id>",
    methods=["POST"]
)
def dispense(queue_id):

    patient = get_queue_patient(
        queue_id
    )


    if patient is None:

        return redirect(
            url_for(
                "prescription_result"
            )
        )


    # ========================================================
    # DAYS SUPPLY CHECK FIRST
    # ========================================================
    # ถ้าจำนวนยาไม่สัมพันธ์กับวันที่จ่าย -> วันนัด
    # ให้เข้าสู่ Consult ก่อน เพื่อให้เห็นว่าขาด/เกินกี่เม็ด

    if patient.get("days_check_required", False):

        patient = ensure_days_consult_pdf(
            patient,
            queue_id=queue_id
        )

        if not patient.get("days_consult_pdf"):
            return render_template(
                "prescription_result.html",
                patients=get_pending_patients(),
                error=(
                    "ไม่สามารถสร้าง Consult PDF "
                    "สำหรับตรวจสอบจำนวนยาและจำนวนวันได้"
                )
            )

        return redirect(
            url_for(
                "doctor_confirm",
                queue_id=queue_id
            )
        )


    # ========================================================
    # INTERACTION CHECK
    # ========================================================

    interaction_results = patient.get(
        "interaction_results",
        []
    )

    if len(interaction_results) > 0:
        return redirect(
            url_for(
                "doctor_confirm",
                queue_id=queue_id
            )
        )


    # ========================================================
    # STOCK CHECK
    # ========================================================

    if not patient_stock_is_sufficient(
        patient
    ):

        return render_template(
            "prescription_result.html",
            patients=get_pending_patients(),
            error=(
                "ไม่สามารถจ่ายยาได้ "
                "เนื่องจาก Stock ไม่เพียงพอ "
                "หรือไม่พบข้อมูลยาใน stock.xlsx"
            )
        )


    # ========================================================
    # SAVE HISTORY + DEDUCT STOCK + RESOLVE
    # ========================================================

    success = complete_queue_action(

        patient,

        "dispense",

        "จ่ายยา",

        queue_id=queue_id

    )


    if not success:

        return render_template(
            "prescription_result.html",
            patients=get_pending_patients(),
            error=(
                "ไม่สามารถบันทึกการจ่ายยาได้ "
                "หรือไม่สามารถหัก Stock ได้ "
                "กรุณาตรวจสอบ stock.xlsx แล้วลองใหม่อีกครั้ง"
            )
        )


    return redirect(
        url_for(
            "prescription_result"
        )
    )


# ============================================================
# DOCTOR CONFIRM
# ============================================================

@app.route(
    "/doctor-confirm/<int:queue_id>",
    methods=["GET"]
)
def doctor_confirm(queue_id):

    patient = get_queue_patient(
        queue_id
    )


    if patient is None:

        return redirect(
            url_for(
                "prescription_result"
            )
        )


    has_interaction = len(
        patient.get(
            "interaction_results",
            []
        )
    ) > 0

    has_days_problem = bool(
        patient.get(
            "days_check_required",
            False
        )
    )

    if not has_interaction and not has_days_problem:

        return redirect(
            url_for(
                "prescription_result"
            )
        )


    if has_days_problem and not has_interaction:
        return render_template(
            "doctor_days_confirm.html",
            patient=patient,
            days_check=patient.get(
                "days_check_results",
                []
            ),
            days_consult_pdf=patient.get(
                "days_consult_pdf",
                ""
            ),
            has_days_problem=has_days_problem
        )

    return render_template(
        "doctor_confirm.html",
        patient=patient,
        substitute_drugs=TEST_SUBSTITUTE_DRUGS,
        show_substitute=False,
        days_check=patient.get("days_check_results", []),
        days_consult_pdf=patient.get("days_consult_pdf", ""),
        has_days_problem=has_days_problem
    )


# ============================================================
# DOCTOR DECISION
# ============================================================

@app.route(
    "/doctor-decision/<int:queue_id>",
    methods=["POST"]
)
def doctor_decision(queue_id):

    patient = get_queue_patient(
        queue_id
    )


    if patient is None:

        return redirect(
            url_for(
                "prescription_result"
            )
        )


    has_interaction = len(
        patient.get(
            "interaction_results",
            []
        )
    ) > 0

    has_days_problem = bool(
        patient.get(
            "days_check_required",
            False
        )
    )

    if not has_interaction and not has_days_problem:

        return redirect(
            url_for(
                "prescription_result"
            )
        )


    # ========================================================
    # GET DECISION
    # ========================================================
    # รองรับค่าจากปุ่มยืนยัน/ยกเลิกหลายรูปแบบ
    decision = (
        request.form.get("decision")
        or request.form.get("action")
        or request.form.get("choice")
        or request.form.get("confirm")
        or ""
    )

    decision = str(decision).strip().lower()

    if decision in {
        "original", "confirm", "confirmed", "yes",
        "approve", "approved", "confirm_original", "use_original",
        "confirm_dispense", "dispense_confirm", "confirm_original_drug"
    }:
        decision = "original"
    elif decision in {
        "override_doctor", "confirm_override",
        "ยืนยันแทนแพทย์", "override"
    }:
        decision = "override_doctor"
    elif decision in {
        "increase", "add", "เพิ่ม",
        "increase_quantity", "add_quantity"
    }:
        decision = "increase"
    elif decision in {
        "decrease", "reduce", "ลด",
        "decrease_quantity", "reduce_quantity"
    }:
        decision = "decrease"
    elif decision in {
        "correct", "corrected", "ถูกแล้ว",
        "right", "keep", "keep_original"
    }:
        decision = "correct"
    elif decision in {
        "cancel", "cancelled", "canceled", "no",
        "reject", "rejected", "cancel_dispense", "cancel_confirm"
    }:
        decision = "cancel"

    # ถ้าเป็น Consult เรื่องจำนวนยา/จำนวนวันเพียงอย่างเดียว
    # รองรับปุ่ม "ยืนยันแทนแพทย์" โดยใช้จำนวนยาตามใบสั่งเดิม
    # และหัก Stock เฉพาะเมื่อกดยืนยันสำเร็จ
    if (
        has_days_problem
        and not has_interaction
        and decision in {
            "override_doctor",
            "confirm_override",
            "ยืนยันแทนแพทย์"
        }
    ):
        decision = "override_doctor"

    # ========================================================
    # CONFIRM INSTEAD OF DOCTOR
    # ========================================================
    # ใช้จำนวนที่ระบุในใบสั่งแพทย์เดิมเท่านั้น
    # ไม่เปลี่ยน quantity เป็นค่าที่ระบบคำนวณ
    if (
        has_days_problem
        and not has_interaction
        and decision == "override_doctor"
    ):
        if not patient_stock_is_sufficient(patient):
            return render_template(
                "doctor_days_confirm.html",
                patient=patient,
                days_check=patient.get("days_check_results", []),
                days_consult_pdf=patient.get("days_consult_pdf", ""),
                error=(
                    "ไม่สามารถยืนยันแทนแพทย์ได้ เนื่องจาก Stock ไม่เพียงพอ "
                    "หรือไม่พบข้อมูลยาใน stock.xlsx"
                )
            )

        patient["_manual_days_decision"] = "override_doctor"
        patient["days_check_required"] = False

        success = complete_queue_action(
            patient,
            "override_doctor",
            "ยืนยันแทนแพทย์ตามใบสั่งยาเดิม",
            queue_id=queue_id
        )

        if not success:
            return render_template(
                "doctor_days_confirm.html",
                patient=patient,
                days_check=patient.get("days_check_results", []),
                days_consult_pdf=patient.get("days_consult_pdf", ""),
                error=(
                    "ไม่สามารถบันทึกการยืนยันแทนแพทย์ได้ "
                    "หรือไม่สามารถหัก Stock ได้"
                )
            )

        return redirect(url_for("prescription_result"))


    # ========================================================
    # DAYS SUPPLY DECISION
    # ========================================================

    if has_days_problem and decision in {
        "increase",
        "decrease",
        "correct",
        "override_doctor"
    }:

        # ----------------------------------------------------
        # ยืนยันแทนแพทย์ = ใช้จำนวนตามใบสั่งเดิม
        # ----------------------------------------------------
        if decision == "override_doctor":

            if not patient_stock_is_sufficient(patient):
                return render_template(
                    "doctor_days_confirm.html",
                    patient=patient,
                    days_check=patient.get("days_check_results", []),
                    days_consult_pdf=patient.get("days_consult_pdf", ""),
                    error=(
                        "ไม่สามารถยืนยันแทนแพทย์ได้ เนื่องจาก Stock ไม่เพียงพอ "
                        "หรือไม่พบข้อมูลยาใน stock.xlsx"
                    )
                )

            patient["_manual_days_decision"] = "override_doctor"
            patient["days_check_required"] = False

            success = complete_queue_action(
                patient,
                "override_doctor",
                "ยืนยันแทนแพทย์ตามใบสั่งยาเดิม",
                queue_id=queue_id
            )

            if not success:
                return render_template(
                    "doctor_days_confirm.html",
                    patient=patient,
                    days_check=patient.get("days_check_results", []),
                    days_consult_pdf=patient.get("days_consult_pdf", ""),
                    error=(
                        "ไม่สามารถบันทึกการยืนยันแทนแพทย์ได้ "
                        "หรือไม่สามารถหัก Stock ได้"
                    )
                )

            return redirect(url_for("prescription_result"))

        # ----------------------------------------------------
        # เพิ่ม / ลด = ปรับจำนวนยาตามจำนวนที่ควรเป็น
        # ----------------------------------------------------
        if decision in {"increase", "decrease"}:

            for medicine in patient.get("medicines", []):

                # ปรับเฉพาะรายการที่ไม่สัมพันธ์กับวันนัด
                if medicine.get("days_match", False):
                    continue

                expected = medicine.get("expected_quantity", 0)

                try:
                    expected = float(expected)
                except Exception:
                    continue

                if expected.is_integer():
                    expected = int(expected)

                medicine["quantity"] = expected

                try:
                    times = float(
                        medicine.get("times_per_day", 1) or 1
                    )
                    medicine["days_supply"] = (
                        float(expected) / times
                        if times > 0 else 0
                    )
                except Exception:
                    medicine["days_supply"] = 0

                medicine["missing_quantity"] = 0
                medicine["excess_quantity"] = 0
                medicine["status"] = "พอดี"
                medicine["days_match"] = True

            patient["_manual_days_decision"] = decision

        else:
            # ถูกแล้ว = แพทย์ยืนยันใช้จำนวนเดิม
            patient["_manual_days_decision"] = "correct"

        # ผ่าน Consult แล้ว แต่ยังไม่หัก Stock
        patient["days_check_required"] = False

        if not update_queue_patient(queue_id, patient):
            return render_template(
                "doctor_days_confirm.html",
                patient=patient,
                days_check=patient.get("days_check_results", []),
                days_consult_pdf=patient.get("days_consult_pdf", ""),
                error="ไม่สามารถบันทึกผลการตรวจจำนวนยาได้ กรุณาลองใหม่"
            )

        return redirect(url_for("prescription_result"))


    # ========================================================
    # ORIGINAL / INTERACTION DECISION
    # ========================================================

    if decision == "original":

        if not patient_stock_is_sufficient(patient):
            return render_template(
                "doctor_confirm.html",
                patient=patient,
                substitute_drugs=TEST_SUBSTITUTE_DRUGS,
                show_substitute=False,
                days_check=patient.get("days_check_results", []),
                days_consult_pdf=patient.get("days_consult_pdf", ""),
                has_days_problem=has_days_problem,
                error=(
                    "ไม่สามารถยืนยันจ่ายยาเดิมได้ เนื่องจาก Stock ไม่เพียงพอ "
                    "หรือไม่พบข้อมูลยาใน stock.xlsx"
                )
            )

        success = complete_queue_action(
            patient,
            "original",
            "ยืนยันใช้ยาเดิม",
            queue_id=queue_id
        )

        if not success:
            return render_template(
                "doctor_confirm.html",
                patient=patient,
                substitute_drugs=TEST_SUBSTITUTE_DRUGS,
                show_substitute=False,
                days_check=patient.get("days_check_results", []),
                days_consult_pdf=patient.get("days_consult_pdf", ""),
                has_days_problem=has_days_problem,
                error=(
                    "ไม่สามารถบันทึกผลการตัดสินใจได้ "
                    "หรือไม่สามารถหัก Stock ได้ "
                    "กรุณาตรวจสอบ stock.xlsx แล้วลองใหม่อีกครั้ง"
                )
            )

        return redirect(url_for("prescription_result"))


    # ========================================================
    # SUBSTITUTE
    # ========================================================

    if decision == "substitute":

        return render_template(

            "doctor_confirm.html",

            patient=patient,

            substitute_drugs=
                TEST_SUBSTITUTE_DRUGS,

            show_substitute=True

        )


    # ========================================================
    # CANCEL
    # ========================================================

    if decision == "cancel":

        success = complete_queue_action(

            patient,

            "cancel",

            "ยกเลิกการจ่ายยา",

            queue_id=queue_id

        )


        if not success:

            return render_template(
                "doctor_confirm.html",

                patient=patient,

                substitute_drugs=
                    TEST_SUBSTITUTE_DRUGS,

                show_substitute=False,

                error=(
                    "ไม่สามารถบันทึกการยกเลิกได้ "
                    "กรุณาลองใหม่อีกครั้ง"
                )
            )


        return redirect(
            url_for(
                "prescription_result"
            )
        )


    return render_template(

        "doctor_confirm.html",

        patient=patient,

        substitute_drugs=
            TEST_SUBSTITUTE_DRUGS,

        show_substitute=False,

        error="กรุณาเลือกการดำเนินการ"

    )


# ============================================================
# CONFIRM SUBSTITUTE
# ============================================================

@app.route(
    "/confirm-substitute/<int:queue_id>",
    methods=["POST"]
)
def confirm_substitute(queue_id):

    patient = get_queue_patient(
        queue_id
    )


    if patient is None:

        return redirect(
            url_for(
                "prescription_result"
            )
        )


    substitute_drug = request.form.get(
        "substitute_drug",
        ""
    ).strip()


    if not substitute_drug:

        return render_template(

            "doctor_confirm.html",

            patient=patient,

            substitute_drugs=
                TEST_SUBSTITUTE_DRUGS,

            show_substitute=True,

            error="กรุณาเลือกยาทดแทน"

        )


    valid_substitute = None


    for drug in TEST_SUBSTITUTE_DRUGS:

        if drug["name"] == substitute_drug:

            valid_substitute = drug

            break


    if valid_substitute is None:

        return render_template(

            "doctor_confirm.html",

            patient=patient,

            substitute_drugs=
                TEST_SUBSTITUTE_DRUGS,

            show_substitute=True,

            error="ไม่พบยาทดแทนในรายการทดสอบ"

        )


    # ========================================================
    # SUBSTITUTE
    #
    # ปัจจุบันยังไม่หัก Stock
    # เพราะ TEST_SUBSTITUTE_DRUGS
    # ไม่มีจำนวนที่ต้องจ่าย
    # ========================================================

    success = complete_queue_action(

        patient,

        "substitute",

        "เลือกยาทดแทน",

        substitute_drug,

        queue_id=queue_id

    )


    if not success:

        return render_template(

            "doctor_confirm.html",

            patient=patient,

            substitute_drugs=
                TEST_SUBSTITUTE_DRUGS,

            show_substitute=True,

            error=(
                "ไม่สามารถบันทึกยาทดแทนได้ "
                "กรุณาลองใหม่อีกครั้ง"
            )

        )


    return redirect(
        url_for(
            "prescription_result"
        )
    )


# ============================================================
# AUTO REFRESH PRESCRIPTION QUEUE
# ============================================================

@app.after_request
def inject_prescription_live_refresh(response):
    """
    หน้าใบสั่งยาจะตรวจ Excel อัตโนมัติทุก 3 วินาที
    และแสดง Days Supply Alert บนหน้าเว็บโดยไม่ต้องแก้ template
    """
    try:
        if (
            request.path == "/prescription-result"
            and response.content_type.startswith("text/html")
            and response.status_code == 200
        ):
            html = response.get_data(as_text=True)

            # ------------------------------------------------
            # DAYS SUPPLY ALERT
            # ------------------------------------------------
            try:
                patients = get_pending_patients()
                alert_blocks = []

                for patient in patients:
                    problems = [
                        item for item in patient.get("days_check_results", [])
                        if item.get("decision") in {
                            "increase", "decrease", "consult"
                        }
                    ]

                    if not problems:
                        continue

                    hn = html_escape(patient.get("hn", ""))
                    name = html_escape(patient.get("name", "ไม่ระบุชื่อ"))
                    dispense_date = html_escape(
                        patient.get("dispense_date", "ไม่ระบุ")
                    )
                    appointment_date = html_escape(
                        patient.get("appointment_date", "ไม่ระบุ")
                    )

                    rows = []
                    for item in problems:
                        drug = html_escape(item.get("name", ""))
                        qty = html_escape(format_quantity(item.get("quantity", 0)))
                        expected = html_escape(
                            format_quantity(item.get("expected_quantity", 0))
                        )
                        missing = item.get("missing_quantity", 0) or 0
                        excess = item.get("excess_quantity", 0) or 0

                        if float(missing) > 0:
                            message = "ขาด {} เม็ด".format(
                                html_escape(format_quantity(missing))
                            )
                        elif float(excess) > 0:
                            message = "เกิน {} เม็ด".format(
                                html_escape(format_quantity(excess))
                            )
                        else:
                            message = html_escape(item.get("status", "คำนวณไม่ได้"))

                        rows.append(
                            "<li><strong>{}</strong> — สั่ง {} เม็ด / ควร {} เม็ด → <strong>{}</strong></li>".format(
                                drug, qty, expected, message
                            )
                        )

                    pdf_filename = patient.get("days_consult_pdf", "")
                    pdf_link = ""
                    if pdf_filename:
                        safe_filename = html_escape(str(pdf_filename))
                        pdf_link = (\
                            '<a href="{}" target="_blank" '
                            'style="display:inline-block;margin-top:10px;padding:10px 16px;'
                            'background:#6f42c1;color:white;text-decoration:none;border-radius:8px;font-weight:bold;">'
                            '📄 ดาวน์โหลด Consult ส่งแพทย์</a>'
                        ).format(
                            url_for("download_consult", filename=safe_filename)
                        )

                    doctor_link = (\
                        '<a href="{}" '
                        'style="display:inline-block;margin-top:10px;margin-left:8px;padding:10px 16px;'
                        'background:#7b1fa2;color:white;text-decoration:none;border-radius:8px;font-weight:bold;">'
                        '👨‍⚕️ เข้าสู่หน้า Consult</a>'
                    ).format(
                        url_for("doctor_confirm", queue_id=patient.get("queue_id"))
                    )

                    alert_blocks.append(
                        """
                        <div style="background:#fff3f3;border:2px solid #e53935;border-radius:12px;padding:18px;margin:0 0 20px 0;box-shadow:0 2px 8px rgba(0,0,0,.05);">
                            <div style="font-size:20px;font-weight:bold;color:#b71c1c;margin-bottom:8px;">🚨 ตรวจพบจำนวนยาไม่สัมพันธ์กับวันนัด</div>
                            <div style="line-height:1.8;">
                                <strong>HN:</strong> {hn} &nbsp; <strong>ผู้ป่วย:</strong> {name}<br>
                                <strong>วันที่จ่าย:</strong> {dispense_date} &nbsp; <strong>วันนัด:</strong> {appointment_date}<br>
                                <strong>จำนวนวันที่ต้องใช้:</strong> {days} วัน
                                <ul style="margin:8px 0 0 20px;">{rows}</ul>
                                {pdf_link}{doctor_link}
                            </div>
                        </div>
                        """.format(
                            hn=hn,
                            name=name,
                            dispense_date=dispense_date,
                            appointment_date=appointment_date,
                            days=html_escape(format_quantity(patient.get("required_days", 0))),
                            rows="".join(rows),
                            pdf_link=pdf_link,
                            doctor_link=doctor_link
                        )
                    )

                if alert_blocks:
                    alert_html = (
                        '<div id="days-supply-alerts" style="margin:0 0 25px 0;">'
                        + "".join(alert_blocks)
                        + "</div>"
                    )

                    lower = html.lower()
                    marker = '<h1>📋 ผลการตรวจใบสั่งยา</h1>'
                    if marker.lower() in lower:
                        pos = lower.find(marker.lower()) + len(marker)
                        html = html[:pos] + alert_html + html[pos:]
                    elif "</body>" in lower:
                        pos = lower.rfind("</body>")
                        html = html[:pos] + alert_html + html[pos:]
            except Exception as alert_error:
                print(
                    "ERROR INJECT DAYS SUPPLY ALERT:",
                    repr(alert_error)
                )

            # ------------------------------------------------
            # REAL-TIME REFRESH
            # ------------------------------------------------
            script = """
<script>
(function () {
    let lastSignature = null;
    let firstCheck = true;

    async function checkPrescriptionLive() {
        try {
            const response = await fetch(
                "/prescription-live",
                {
                    cache: "no-store",
                    headers: {"X-Requested-With": "XMLHttpRequest"}
                }
            );

            if (!response.ok) return;

            const data = await response.json();

            if (firstCheck) {
                lastSignature = data.signature;
                firstCheck = false;
                return;
            }

            if (lastSignature !== data.signature) {
                window.location.reload();
            }
        } catch (error) {
            console.log("Prescription live refresh:", error);
        }
    }

    checkPrescriptionLive();
    setInterval(checkPrescriptionLive, 3000);
})();
</script>
"""

            lower = html.lower()
            if "</body>" in lower:
                position = lower.rfind("</body>")
                html = html[:position] + script + html[position:]
            else:
                html += script

            response.set_data(html)

    except Exception as e:
        print(
            "ERROR INJECT PRESCRIPTION LIVE REFRESH:",
            repr(e)
        )

    return response


# ============================================================
# HTML ESCAPE
# ============================================================

def html_escape(value):
    """Escape text before injecting dynamic values into HTML."""
    import html
    return html.escape(str(value if value is not None else ""), quote=True)


# ============================================================
# HOME
# ============================================================

@app.route("/")
def index():

    return render_template(
        "index.html"
    )


# ============================================================
# STOCK
# ============================================================

@app.route(
    "/stock",
    methods=["GET", "POST"]
)
def stock():

    # ========================================================
    # GET
    # ========================================================

    if request.method == "GET":

        stock_data = load_stock()


        stock_exists = os.path.exists(
            STOCK_FILE
        )


        stock_filename = (

            os.path.basename(
                STOCK_FILE
            )

            if stock_exists

            else None

        )


        stock_updated_at = None


        if stock_exists:

            try:

                modified_time = os.path.getmtime(
                    STOCK_FILE
                )


                stock_updated_at = datetime.fromtimestamp(
                    modified_time
                ).strftime(
                    "%d/%m/%Y %H:%M:%S"
                )

            except Exception:

                stock_updated_at = None


        return render_template(

            "stock.html",

            stock_data=stock_data,

            stock_exists=stock_exists,

            stock_filename=stock_filename,

            stock_updated_at=stock_updated_at

        )


    # ========================================================
    # POST
    # ========================================================

    uploaded_file = request.files.get(
        "stock_file"
    )


    if uploaded_file is None:

        return render_template(

            "stock.html",

            stock_data=load_stock(),

            stock_exists=os.path.exists(
                STOCK_FILE
            ),

            stock_filename=(

                os.path.basename(
                    STOCK_FILE
                )

                if os.path.exists(
                    STOCK_FILE
                )

                else None

            ),

            error="กรุณาเลือกไฟล์ Excel Stock"

        )


    if uploaded_file.filename == "":

        return render_template(

            "stock.html",

            stock_data=load_stock(),

            stock_exists=os.path.exists(
                STOCK_FILE
            ),

            stock_filename=(

                os.path.basename(
                    STOCK_FILE
                )

                if os.path.exists(
                    STOCK_FILE
                )

                else None

            ),

            error="กรุณาเลือกไฟล์ Excel Stock"

        )


    extension = os.path.splitext(
        uploaded_file.filename
    )[1].lower()


    # ========================================================
    # CHECK EXTENSION
    # ========================================================

    if extension not in [
        ".xlsx",
        ".xlsm"
    ]:

        return render_template(

            "stock.html",

            stock_data=load_stock(),

            stock_exists=os.path.exists(
                STOCK_FILE
            ),

            stock_filename=(

                os.path.basename(
                    STOCK_FILE
                )

                if os.path.exists(
                    STOCK_FILE
                )

                else None

            ),

            error=(
                "กรุณาเลือกไฟล์ Excel "
                "(.xlsx หรือ .xlsm)"
            )

        )


    temp_stock_path = None


    try:

        # ====================================================
        # SAVE UPLOAD TO TEMP
        # ====================================================

        with tempfile.NamedTemporaryFile(
            delete=False,
            suffix=extension
        ) as temp_file:

            uploaded_file.save(
                temp_file.name
            )

            temp_stock_path = temp_file.name


        # ====================================================
        # TEST OPEN EXCEL
        # ====================================================

        wb = load_workbook(
            temp_stock_path,
            data_only=True,
            read_only=True
        )


        ws = wb.active


        # ====================================================
        # CHECK SHEET
        # ====================================================

        if ws.max_row < 1:

            wb.close()

            raise ValueError(
                "ไฟล์ Excel ไม่มีข้อมูล"
            )


        headers = [

            cell.value

            for cell in ws[1]

        ]


        # ====================================================
        # CHECK HEADER
        # ====================================================

        if not any(
            header is not None
            for header in headers
        ):

            wb.close()

            raise ValueError(
                "ไม่พบหัวตารางในไฟล์ Stock"
            )


        # ====================================================
        # CHECK DATA
        # ====================================================

        has_data = False


        for row in ws.iter_rows(
            min_row=2,
            values_only=True
        ):

            if any(
                value is not None
                for value in row
            ):

                has_data = True

                break


        wb.close()


        if not has_data:

            raise ValueError(
                "ไม่พบข้อมูล Stock ใต้หัวตาราง"
            )


        # ====================================================
        # REPLACE OLD STOCK
        # ====================================================

        shutil.copyfile(
            temp_stock_path,
            STOCK_FILE
        )


        print("")
        print("=" * 70)
        print("STOCK FILE UPDATED")
        print(
            "Original uploaded filename:",
            uploaded_file.filename
        )
        print(
            "Stored as:",
            STOCK_FILE
        )
        print("=" * 70)


        # ====================================================
        # LOAD NEW STOCK
        # ====================================================

        new_stock_data = load_stock()


        # ====================================================
        # CHECK AGAIN
        # ====================================================

        if not new_stock_data:

            raise ValueError(
                "อัปโหลดสำเร็จแต่ไม่สามารถอ่านข้อมูล Stock ได้"
            )


        return render_template(

            "stock.html",

            stock_data=new_stock_data,

            stock_exists=True,

            stock_filename=os.path.basename(
                STOCK_FILE
            ),

            stock_updated_at=datetime.now().strftime(
                "%d/%m/%Y %H:%M:%S"
            ),

            success=(
                "อัปโหลดไฟล์ Stock สำเร็จ "
                "ระบบเปลี่ยนเป็นข้อมูลจากไฟล์ใหม่แล้ว"
            )

        )


    except Exception as e:

        print("")
        print("=" * 70)
        print("ERROR UPLOAD STOCK")
        print(repr(e))
        print("=" * 70)


        return render_template(

            "stock.html",

            stock_data=load_stock(),

            stock_exists=os.path.exists(
                STOCK_FILE
            ),

            stock_filename=(

                os.path.basename(
                    STOCK_FILE
                )

                if os.path.exists(
                    STOCK_FILE
                )

                else None

            ),

            error=(
                "ไม่สามารถอัปโหลดไฟล์ Stock ได้: "
                + str(e)
            )

        )


    finally:

        if temp_stock_path:

            try:

                if os.path.exists(
                    temp_stock_path
                ):

                    os.remove(
                        temp_stock_path
                    )

            except Exception:

                pass


# ============================================================
# INTERACTIONS
# ============================================================

@app.route("/interactions")
def interactions():

    data = load_interactions()


    return render_template(
        "interactions.html",
        interactions=data
    )


# ============================================================
# INTERACTION CHECK
# ============================================================

@app.route(
    "/interaction-check",
    methods=["GET", "POST"]
)
def interaction_check():

    results = []

    drug1 = ""

    drug2 = ""


    if request.method == "POST":

        drug1 = request.form.get(
            "drug1",
            ""
        )


        drug2 = request.form.get(
            "drug2",
            ""
        )


        results = check_drug_interactions(
            [
                drug1,
                drug2
            ]
        )


    return render_template(
        "interaction_check.html",
        results=results,
        drug1=drug1,
        drug2=drug2
    )


# ============================================================
# APPOINTMENT
# ============================================================

@app.route("/appointment")
def appointment():

    return render_template(
        "appointment.html"
    )


# ============================================================
# HISTORY
# ============================================================

@app.route("/history")
def history():

    history_data = get_history()


    return render_template(
        "history.html",
        history=history_data
    )


# ============================================================
# CLEAR HISTORY
# TESTING ONLY
# ============================================================

@app.route(
    "/clear-history",
    methods=["POST"]
)
def clear_history():

    conn = get_db()


    try:

        conn.execute(
            "DELETE FROM medication_history"
        )


        conn.execute(
            """
            DELETE FROM prescription_queue
            WHERE status = 'resolved'
            """
        )


        conn.commit()


    except Exception as e:

        conn.rollback()

        print(
            "ERROR CLEAR HISTORY:",
            repr(e)
        )


    finally:

        conn.close()


    return redirect(
        url_for(
            "history"
        )
    )


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":

    print("=" * 60)


    print(
        "Medication Management System"
    )


    print(
        "Server: http://127.0.0.1:5000"
    )


    print(
        "Prescription Upload:"
    )


    print(
        "http://127.0.0.1:5000/prescription"
    )


    print(
        "Stock:"
    )


    print(
        "http://127.0.0.1:5000/stock"
    )


    print("=" * 60)


    app.run(
        debug=True
    )
